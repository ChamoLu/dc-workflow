# -*- coding: utf-8 -*-
r"""
采集数据源：中华人民共和国宁波海关
数据来源编码：1020000624
数据源路径代码：2076929850064633857
路径：中华人民共和国宁波海关\政府信息公开\法定主动公开内容\统计信息\统计数据\各大洲进出口情况表
老采集站点名称：地方经济运行-M1-宁波市-宁波市海关_各大洲
文件名称：2026年4月宁波口岸分洲别进出口总值表.xlsx
目标表：usrDFJJYXSJB
excelrule行号：536
源文件：<运行时通过 fileMd5 下载，不使用本地路径>
指标名称示例：亚洲进出口本月

拆解规则：
1. 指标名称取值逻辑：A列+行4行5
2. 指标名称标准化处理：
3. 指标单位取值逻辑：文件中的进出口、出口、进口绝对量统一取 万元；“累计与去年同期±%”下的进出口、出口、进口同比指标取 %。
4. 指标单位位置：绝对量单位取表格右上角 “单位：万元”；同比单位取第4行表头 “累计与去年同期±%”。
5. 截止日期取值逻辑：文件为2026年4月数据：“本月”指标取 2026-04-30；“一至本月”累计及累计同比同样截止 2026-04-30。
6. 截止日期位置：年月取标题 “2026年4月宁波口岸分洲别进出口总值表”，期间类型取第5行 “本月 / 一至本月”。
7. 备注/Sheet限制：
8. 全局规范：全角转半角；去空格、顿号、英文逗号、单引号和换行；不伪造缺失字段。
"""
import calendar
import csv
import io
import json
import re
import unicodedata
import requests

BASE_DOWNLOAD_URL = "http://service-1.example.invalid/api/file/download/"
XLSX_MAGIC = b"PK"
XLS_MAGIC = bytes.fromhex("D0CF11E0A1B11AE1")
CSV_ENCODINGS = ("utf-8-sig", "gbk", "utf-8", "gb2312")
RULE = json.loads(r"""{"row": 536, "kind": "cn_trade", "path_code": "2076929850064633857", "allowed_names": []}""")


def text(value):
    return "" if value is None else str(value).strip()


def clean_name(value):
    value = unicodedata.normalize("NFKC", text(value))
    return re.sub(r"[\s、,'\n]+", "", value)


CANONICAL = {clean_name(name): name for name in RULE.get("allowed_names", [])}


def canonical(name):
    name = clean_name(name)
    return CANONICAL.get(name, name)


def number(value):
    value = text(value).replace(",", "").replace("△", "-")
    if value in ("", "-", "--", "---", "ND", "N/A", "NA", "#N/A"):
        return None
    try:
        float(value)
    except ValueError:
        return None
    return value


def month_end(year, month):
    return f"{int(year):04d}-{int(month):02d}-{calendar.monthrange(int(year), int(month))[1]:02d}"


def period(value, monthly=False, exact=False, fiscal=False):
    year = getattr(value, "year", None)
    month = getattr(value, "month", None)
    day = getattr(value, "day", None)
    if isinstance(year, int) and isinstance(month, int) and isinstance(day, int):
        return f"{year:04d}-{month:02d}-{day:02d}" if exact else month_end(year, month)
    value = text(value)
    if not value:
        return ""
    m = re.search(r"(19|20)\d{2}[-/.](\d{1,2})[-/.](\d{1,2})", value)
    if m:
        return f"{int(m.group(0)[:4]):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}" if exact else month_end(m.group(0)[:4], m.group(2))
    m = re.search(r"((?:19|20)\d{2})[-/\.](\d{1,2})", value)
    if m:
        return month_end(m.group(1), m.group(2))
    m = re.search(r"((?:19|20)\d{2}).*?(?:Q|季度|Quarter)\s*([1-4])", value, re.I)
    if m:
        return month_end(m.group(1), int(m.group(2)) * 3)
    m = re.search(r"((?:19|20)\d{2}).*?(\d{1,2})\s*(?:[-~～至到]\s*\d{1,2})?月", value)
    if m:
        nums = re.findall(r"\d{1,2}", value[m.start():])
        return month_end(m.group(1), nums[-1])
    m = re.search(r"((?:19|20)\d{2})", value)
    if m:
        return f"{int(m.group(1)) + 1:04d}-03-31" if fiscal and re.search(r"FY|F\.Y\.|年度", value, re.I) else f"{m.group(1)}-12-31"
    return ""


def record(data, name, value, unit, end_date):
    value = number(value)
    name = canonical(name)
    if name and value is not None and unit and end_date:
        data.append({"indicator_name": name, "indicator_data": value, "indicator_unit": unit, "end_date": end_date})


def read_csv(raw):
    for encoding in CSV_ENCODINGS:
        try:
            return [("CSV", list(csv.reader(io.StringIO(raw.decode(encoding)))))]
        except UnicodeDecodeError:
            pass
    raise ValueError("CSV 编码无法识别")


def read_xlsx(raw):
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    return [(sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)]) for sheet in book.worksheets]


def read_xls(raw):
    import xlrd
    book = xlrd.open_workbook(file_contents=raw)
    out = []
    for sheet in book.sheets():
        rows = []
        for r in range(sheet.nrows):
            values = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    y, m, d, _, _, _ = xlrd.xldate_as_tuple(cell.value, book.datemode)
                    values.append(f"{y:04d}-{m:02d}-{d:02d}")
                else:
                    values.append(cell.value)
            rows.append(values)
        out.append((sheet.name, rows))
    return out


def read_file(raw):
    if raw.startswith(XLSX_MAGIC):
        return read_xlsx(raw)
    if raw.startswith(XLS_MAGIC):
        return read_xls(raw)
    return read_csv(raw)


def forward(rows, row_index, column):
    for r in range(row_index, -1, -1):
        if column < len(rows[r]) and text(rows[r][column]):
            return text(rows[r][column])
    return ""


def header_from_left(rows, row_index, column):
    for c in range(column, -1, -1):
        if c < len(rows[row_index]) and text(rows[row_index][c]):
            return text(rows[row_index][c])
    return ""


def parse_wide(rows, header_row, start_row, unit, date_column=0, monthly=True, exact=False, name_prefix="", names=None):
    data = []
    headers = rows[header_row] if header_row < len(rows) else []
    for row in rows[start_row:]:
        end_date = period(row[date_column] if date_column < len(row) else "", monthly, exact)
        if not end_date:
            continue
        for col in range(date_column + 1, min(len(row), len(headers))):
            name = (names[col] if names and col in names else text(headers[col]))
            record(data, name_prefix + name, row[col], unit(name, col, row) if callable(unit) else unit, end_date)
    return data


def parse_rows(sheets):
    kind = RULE["kind"]
    if kind == "cn_trade":
        rows = sheets[0][1]; data = []
        for row in rows[5:]:
            for col in range(1, min(len(row), 10)):
                top = header_from_left(rows, 3, col) if col < 7 else forward(rows, 3, 7)
                sub = text(rows[4][col])
                record(data, text(row[0]) + top + sub, row[col], "%" if col >= 7 else "万元", "2026-04-30")
        return data
    if kind == "oecd":
        rows = sheets[0][1]; header = rows[0]; idx = {x: header.index(x) for x in ("REF_AREA", "ACTIVITY", "TRANSFORMATION", "ADJUSTMENT", "TIME_PERIOD", "OBS_VALUE")}
        data=[]
        for row in rows[1:]:
            name = "-".join(text(row[idx[x]]) for x in ("REF_AREA", "ACTIVITY", "TRANSFORMATION", "ADJUSTMENT"))
            record(data, name, row[idx["OBS_VALUE"]], "%" if text(row[idx["TRANSFORMATION"]]) == "GY" else "指数", period(row[idx["TIME_PERIOD"]], True))
        return data
    if kind == "survey":
        return parse_wide(sheets[0][1], 0, 1, "指数", exact=True)
    if kind in ("ipmp", "arg_fx"):
        data=[]
        for sheet, rows in sheets:
            if kind == "ipmp" and sheet == "Ponderadores": continue
            if not rows: continue
            header_index = 0 if kind == "ipmp" else 1
            headers=rows[header_index]
            for row in rows[header_index + 1:]:
                for col in range(1, min(len(row),len(headers))):
                    name = sheet + re.sub(r"\([^)]*\)", "", text(headers[col])) if kind == "ipmp" else sheet + text(headers[col])
                    mon = kind == "arg_fx" and ("prom" in sheet.lower() or "ponder" in sheet.lower())
                    record(data,name,row[col],"%" if "ponder" in sheet.lower() else "指数",period(row[0],mon,not mon))
        return data
    if kind == "sifma":
        rows = next(rows for sheet, rows in sheets if sheet == "Issuance"); data=[]; latest=""
        for row in rows[8:]:
            d=period(row[0] if row else "",True)
            if d and not text(row[0]).upper().startswith("YTD"): latest=max(latest,d)
        for row in rows[8:]:
            label=text(row[0]); d=latest if label.upper().startswith("YTD") else period(label,True)
            if not d: continue
            for col in range(1,len(row)):
                group=forward(rows,6,col); head=text(rows[7][col]) if col<len(rows[7]) else ""
                record(data,group+head,row[col],"%" if any(x in group for x in ("%","Change")) else "十亿美元",d)
        return data
    if kind == "allmonth":
        def unit(name, *_):
            return "百分点" if name == "Spread" else ("%" if name != "NBER_Rec" else "-")
        return parse_wide(sheets[0][1], 0, 1, unit, monthly=True)
    if kind == "cmdi":
        rows=next(rows for sheet,rows in sheets if sheet=="Index Data"); return parse_wide(rows,5,6,"指数",exact=True)
    if kind == "mct":
        rows=sheets[-1][1]
        return parse_wide(rows,5,6,lambda _,c,_r: "%" if c <= 5 else "百分点")
    if kind == "h15":
        return parse_wide(sheets[0][1],5,6,"%")
    if kind == "finra":
        return parse_wide(sheets[0][1],0,1,"百万美元")
    if kind == "z1":
        rows=sheets[0][1]; data=[]
        for row in rows[1:]:
            for col in range(6,len(row)):
                record(data,row[5],row[col],"百万美元",period(rows[0][col]))
        return data
    if kind == "boj_projection":
        rows=sheets[0][1]; data=[]
        for row in rows[12:]:
            label=text(row[2] if len(row)>2 else "")+text(row[3] if len(row)>3 else "")
            if not label: continue
            for col in range(4,min(len(row),7)):
                head=text(rows[11][col]); d="2025-08-31" if "Previous" in head else "2026-08-31"
                record(data,label+head,row[col],"亿日元",d)
        return data
    if kind == "estat_elasticity":
        rows=sheets[0][1]; data=[]
        for row in rows[7:]:
            label=text(row[0])
            if not label: continue
            for col in range(6,len(row)):
                record(data,("二人" if col<9 else "勤劳")+label+text(rows[5][col]),row[col],"-","2025-12-31")
        return data
    if kind == "estat_building":
        data=[]
        for sheet,rows in sheets:
            if sheet not in ("ken-e_uy","ken-e_iy"): continue
            data.extend(parse_wide(rows,7 if sheet=="ken-e_uy" else 9,10 if sheet=="ken-e_uy" else 11,lambda n,*_: "%" if "Change" in n or "%" in n else "千平方米",name_prefix=sheet))
        return data
    if kind == "boj_fx":
        rows=sheets[0][1]; names={c:text(rows[4][c]) for c in range(1,len(rows[4]))}
        return parse_wide(rows,4,9,"日元/美元",exact=True,names=names)
    if kind == "offshore":
        rows=sheets[0][1]; data=[]; year=""
        for row in rows[12:]:
            for value in row[:4]:
                m=re.search(r"(19|20)\d{2}",text(value))
                if m: year=m.group(0)
            for col in range(4,len(row)):
                name="".join(forward(rows,r,col) for r in (4,6,8,11))
                p=text(row[3] if len(row)>3 else "")
                end=period(year+" "+p,True) if year else ""
                record(data,name,row[col],"亿日元",end)
        return data
    return []


def parse_file_bytes(raw):
    return parse_rows(read_file(raw))


def main(params):
    file_md5 = params.get("fileMd5") or params.get("parseFileMd5")
    if not file_md5:
        return {"code": -1, "msg": "缺少 fileMd5", "ExcelData": []}
    try:
        response = requests.get(BASE_DOWNLOAD_URL.rstrip("/") + "/" + str(file_md5), timeout=60)
        response.raise_for_status()
        return {"code": 0, "msg": "success", "ExcelData": parse_file_bytes(response.content)}
    except Exception as exc:
        return {"code": -1, "msg": str(exc), "ExcelData": []}
