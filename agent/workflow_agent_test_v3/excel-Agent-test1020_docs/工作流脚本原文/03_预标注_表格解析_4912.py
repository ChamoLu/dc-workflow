# 数据源：上街区人民政府
# 数据源路径代码：2085711403154546690
# 数据来源编码：01020002879
# 路径：上街区人民政府\政务公开\政府信息公开\法定主动公开内容\统计信息\统计月报
# 目标表：usrDFJJYXSJB
# 运行时从title解析统计期；不读取GZK，生产输出不含indicator_code。

import re
import calendar
import requests
from io import BytesIO
import openpyxl

_DOWNLOAD_URL = "http://service-1.example.invalid/api/file/download/"


def _clean(value):
    if value is None:
        return ""
    text = str(value).replace("\n", "").replace("\r", "").strip()
    text = text.replace("、", "").replace(",", "").replace("'", "")
    text = text.replace("*", "")
    return text


def _valid(value):
    if value is None or isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return True
    text = _clean(value)
    return bool(text) and text not in {"-", "—", "－", "--", "N/A", "NA"}


def _value(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    return _clean(value)


def _date_from_title(title):
    text = _clean(title)
    matches = re.findall(r"(?<!\d)(20\d{2})年\s*(1[0-2]|0?[1-9])月", text)
    if len(matches) != 1:
        return ""
    year, month = int(matches[0][0]), int(matches[0][1])
    return "%04d-%02d-%02d" % (year, month, calendar.monthrange(year, month)[1])


def _download(params):
    file_md5 = params.get("fileMd5")
    if not file_md5:
        raise ValueError("fileMd5 is required")
    response = requests.get(_DOWNLOAD_URL + str(file_md5), timeout=60)
    response.raise_for_status()
    return response.content


def _append(records, evidence, name, data, unit, end_date, name_refs,
            data_ref, unit_refs, date_refs):
    record = {
        "indicator_name": _clean(name),
        "indicator_data": data,
        "indicator_unit": unit or "-",
        "end_date": end_date or ""
    }
    if not record["indicator_name"]:
        return
    index = len(records)
    records.append(record)
    if evidence is not None:
        evidence.append({
            "record_index": index,
            "name_refs": name_refs,
            "data_ref": data_ref,
            "unit_refs": unit_refs,
            "date_refs": date_refs
        })


def _annotation(value):
    if not isinstance(value, str):
        return None
    match = re.fullmatch(r"(?:提高|上升|增加)\s*([0-9]+(?:\.[0-9]+)?)\s*(个百分点|百分点)", value.strip())
    if match:
        return match.group(1), match.group(2)
    return None


def _extract(ws, title, evidence=None):
    records = []
    end_date = _date_from_title(title)
    date_refs = ["runtime:title"] if end_date else []

    # 首块：累计列C、同比列D；A21脚注规定带星指标增长口径为比年初增长。
    for row in range(4, 21):
        label = _clean(ws.cell(row, 1).value)
        if not label or row == 10:
            if row == 10:
                pass
            continue
        if row in (18, 19):
            growth_name = "比年初增长"
        else:
            growth_name = "比去年同期±%"
        base_unit = "万元"
        if row == 12:
            base_unit = "吨标准煤"
        elif row == 13:
            base_unit = "吨标准煤/万元"
        elif row == 20:
            base_unit = "%"
        c = ws.cell(row, 3).value
        if _valid(c):
            _append(records, evidence, label + "累计", _value(c), base_unit,
                    end_date, ["主要经济指标!A%d" % row, "主要经济指标!C3"],
                    "主要经济指标!C%d" % row,
                    ["主要经济指标!C2" if row not in (12, 13, 20) else "主要经济指标!A%d" % row],
                    date_refs)
        d = ws.cell(row, 4).value
        if row == 20:
            parsed = _annotation(d)
            if parsed:
                data, unit = parsed
                _append(records, evidence, label + growth_name, data, unit,
                        end_date, ["主要经济指标!A20", "主要经济指标!A21"],
                        "主要经济指标!D20", ["主要经济指标!D20"], date_refs)
        elif _valid(d):
            _append(records, evidence, label + growth_name, _value(d), "%",
                    end_date, ["主要经济指标!A%d" % row, "主要经济指标!D3", "主要经济指标!A21" if row in (18, 19) else "主要经济指标!D3"],
                    "主要经济指标!D%d" % row, ["主要经济指标!D3"], date_refs)

    # 第二块：规模以上工业增加值。
    second_names = {
        27: "规模以上工业增加值",
        28: "规模以上工业增加值制造业",
        29: "规模以上工业增加值电热气水的生产和供应业",
        30: "规模以上工业增加值非公有制",
        31: "规模以上工业增加值大中型企业",
        32: "规模以上工业增加值高技术产业",
        33: "规模以上工业增加值高技术产业占规上工业比重（%）",
        34: "规模以上工业增加值高新技术产业",
        35: "规模以上工业增加值高新技术产业占规上工业比重（%）",
        36: "规模以上工业增加值战略新兴产业",
        37: "规模以上工业增加值战略新兴产业占规上工业比重（%）",
        38: "规模以上工业增加值工业主导产业合计",
        39: "规模以上工业增加值工业主导产业合计装备制造",
        40: "规模以上工业增加值工业主导产业合计铝冶炼",
        41: "规模以上工业增加值工业主导产业合计非金属矿物制品业"
    }
    for row, base in second_names.items():
        c = ws.cell(row, 3).value
        d = ws.cell(row, 4).value
        c_unit = "%" if row in (33, 35, 37) else "-"
        if _valid(c):
            _append(records, evidence, base + "累计", _value(c), c_unit, end_date,
                    ["主要经济指标!A%d" % row, "主要经济指标!A24", "主要经济指标!C26"],
                    "主要经济指标!C%d" % row,
                    ["主要经济指标!A%d" % row] if c_unit == "%" else [], date_refs)
        if _valid(d):
            _append(records, evidence, base + "比去年同期±%", _value(d), "%", end_date,
                    ["主要经济指标!A%d" % row, "主要经济指标!A24", "主要经济指标!D26"],
                    "主要经济指标!D%d" % row, ["主要经济指标!D26"], date_refs)

    # 第三块：固定资产投资完成额，仅有同比列B。
    third_names = {
        48: "固定资产投资完成额固定资产投资",
        49: "固定资产投资完成额固定资产投资工业投资",
        50: "固定资产投资完成额固定资产投资房地产投资",
        51: "固定资产投资完成额固定资产投资基础设施",
        52: "固定资产投资完成额固定资产投资民间投资",
        54: "固定资产投资完成额固定资产投资按产业类型分第一产业",
        55: "固定资产投资完成额固定资产投资按产业类型分第二产业",
        56: "固定资产投资完成额固定资产投资按产业类型分第三产业",
        58: "固定资产投资完成额固定资产投资按建设性质分新建",
        59: "固定资产投资完成额固定资产投资按建设性质分扩建",
        60: "固定资产投资完成额固定资产投资按建设性质分改建和技术改造"
    }
    for row, name in third_names.items():
        value = ws.cell(row, 2).value
        if _valid(value):
            refs = ["主要经济指标!A%d" % row, "主要经济指标!A45", "主要经济指标!B47"]
            _append(records, evidence, name + "比去年同期±%", _value(value), "%", end_date,
                    refs, "主要经济指标!B%d" % row, ["主要经济指标!B47"], date_refs)

    # 第四块：社会消费品零售总额。
    fourth_names = {
        66: "社会消费品零售总额",
        68: "社会消费品零售总额按规模分限额以上",
        69: "社会消费品零售总额按规模分限额以下",
        71: "社会消费品零售总额按行业分批发业",
        72: "社会消费品零售总额按行业分零售业",
        73: "社会消费品零售总额按行业分住宿业",
        74: "社会消费品零售总额按行业分餐饮业"
    }
    for row, name in fourth_names.items():
        c = ws.cell(row, 3).value
        d = ws.cell(row, 4).value
        if _valid(c):
            _append(records, evidence, name + "累计", _value(c), "万元", end_date,
                    ["主要经济指标!A%d" % row, "主要经济指标!A63", "主要经济指标!C65"],
                    "主要经济指标!C%d" % row, ["主要经济指标!C64"], date_refs)
        if _valid(d):
            _append(records, evidence, name + "比去年同期±%", _value(d), "%", end_date,
                    ["主要经济指标!A%d" % row, "主要经济指标!A63", "主要经济指标!D65"],
                    "主要经济指标!D%d" % row, ["主要经济指标!D65"], date_refs)
    return records, evidence


def _run(params, with_evidence=False):
    raw = _download(params)
    workbook = openpyxl.load_workbook(BytesIO(raw), data_only=True, read_only=False)
    if "主要经济指标" not in workbook.sheetnames:
        raise ValueError("missing sheet: 主要经济指标")
    records, evidence = _extract(workbook["主要经济指标"], params.get("title", ""), [] if with_evidence else None)
    result = {"code": 0, "msg": "success", "ExcelData": records}
    if with_evidence:
        result["_dc_evidence"] = evidence
    return result


def main(params):
    return _run(params, False)


def _dc_validate(params):
    return _run(params, True)
