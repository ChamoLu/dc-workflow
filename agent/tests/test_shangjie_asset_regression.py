"""Actual March/June bytes: headers, units, date, and complete numeric-cell coverage."""

import calendar
import importlib.util
import json
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

from dc_agent.validation.validator import number

ROOT = Path(__file__).resolve().parents[1]
NAME = "上街区人民政府_政务公开_政府信息公开_法定主动公开内容_统计信息_统计月报.py"
ASSET_PATH = ROOT / "code/excel/xlsx" / NAME


@pytest.fixture(params=[3, 6], ids=["March", "June"])
def source_case(request, monkeypatch):
    month = request.param
    path = ASSET_PATH
    spec = importlib.util.spec_from_file_location("shangjie_original_source_test", path)
    parser = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(parser)
    raw = (ROOT / "tests/fixtures/shangjie" / ("2026-%02d.xlsx" % month)).read_bytes()
    params = {
        "fileMd5": "actual-fixture-%d" % month,
        "title": "上街经济动态2026年%d月*地方经济" % month,
    }

    class Response:
        content = raw

        def raise_for_status(self):
            pass

    def download(url, **kwargs):
        assert url == "http://service-1.example.invalid/api/file/download/" + params["fileMd5"]
        return Response()

    monkeypatch.setattr(parser.requests, "get", download)
    result = parser.main(params)
    traced = parser._dc_validate(params)
    assert {key: value for key, value in traced.items() if key != "_dc_evidence"} == result
    records = result["ExcelData"]
    refs = {
        entry["data_ref"].split("!")[1]: records[entry["record_index"]]
        for entry in traced["_dc_evidence"]
    }
    return month, raw, records, refs


def test_real_sources_have_exact_numeric_coverage_and_no_header_records(source_case):
    month, raw, records, refs = source_case
    assert len(records) == (67 if month == 3 else 70)
    book = openpyxl.load_workbook(BytesIO(raw), data_only=True)
    try:
        sheet = book["主要经济指标"]
        ranges = (
            [(4, 20, (3, 4)), (27, 41, (3, 4)), (48, 60, (2,)), (66, 74, (3, 4))]
            if month == 3
            else [(4, 21, (3, 4)), (28, 42, (3, 4)), (49, 61, (2,)), (67, 75, (3, 4))]
        )
        expected = {
            sheet.cell(row, col).coordinate
            for start, end, columns in ranges
            for row in range(start, end + 1)
            for col in columns
            if number(sheet.cell(row, col).value) is not None
        }
        expected.add("D20" if month == 3 else "D21")  # Actual CPI percentage-point text.
        assert set(refs) == expected and len(refs) == len(records)
        for record in records:
            assert set(record) == {"indicator_name", "indicator_data", "indicator_unit", "end_date"}
            assert number(record["indicator_data"]) is not None
            assert record["end_date"] == "2026-%02d-%02d" % (
                month,
                calendar.monthrange(2026, month)[1],
            )
        json.dumps(records, ensure_ascii=False, allow_nan=False)
    finally:
        book.close()


def test_real_units_parent_labels_and_statistical_period_remain_source_bound(source_case):
    month, _, records, refs = source_case
    cpi_row = 20 if month == 3 else 21
    assert refs["C%d" % cpi_row]["indicator_unit"] == "%"
    assert refs["D%d" % cpi_row]["indicator_unit"] == "个百分点"
    assert refs["D%d" % cpi_row]["indicator_data"] == ("1.2" if month == 3 else "1.4")
    assert refs["C12"]["indicator_unit"] == "吨标准煤"
    assert refs["D13"]["indicator_unit"] == "%"
    assert refs["D18"]["indicator_name"].endswith("比年初增长")
    assert refs["C5"]["indicator_name"] == "第一产业累计"
    assert any("按行业分(按行业销售额划分)批发业" in row["indicator_name"] for row in records)
    assert any("高技术产业占规上工业比重(%)累计" in row["indicator_name"] for row in records)
    if month == 6:
        assert refs["C20"]["indicator_unit"] == "元"
        assert refs["D20"]["indicator_unit"] == "%"
        assert refs["C20"]["indicator_data"] == 28716.29
        assert "C27" not in refs and "D27" not in refs
