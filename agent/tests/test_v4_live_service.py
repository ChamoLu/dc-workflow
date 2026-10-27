"""Real V4 loopback HTTP, Windows spawn worker, and delivered platform code.

The three assets are deliberately hand-written, reviewed transport fixtures in an
isolated registered store. They never read production assets or call a model or
the company network. Excel and HTML platform execution reads controlled source
bytes; Markdown delivery does not claim target-model extraction accuracy.
"""

import hashlib
import importlib.util
import json
import multiprocessing
import shutil
import socket
import threading
import time
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import openpyxl
import pytest
import requests
import uvicorn

from dc_agent.assets.store import ArtifactStore
from dc_agent.config import Settings
from dc_agent.domain import Task
from dc_agent.entrypoints.http import create_app
from dc_agent.infrastructure.storage import Repository
from dc_agent.integration.worker import terminate_worker_tree

ROOT = Path(__file__).resolve().parents[1]
NODES = ROOT / "workflow_agent_v4" / "工作流节点代码"
TOKEN = "v4-live-controlled-fixture-token-not-a-production-credential"
DOWNLOAD = "http://service-1.example.invalid/api/file/download/"
FIELDS = (
    "status",
    "stage",
    "job_id",
    "PythonFile",
    "PythonCode",
    "code_sha256",
    "artifact_version",
    "cache_hit",
    "quality_status",
    "warning_summary",
    "error_code",
    "message",
    "report_ref",
    "elapsed_seconds",
    "PromptFile",
    "PromptCode",
)
RECORD = {
    "indicator_name": "传输夹具指标(万元)",
    "indicator_data": "0.0000",
    "indicator_unit": "万元",
    "end_date": "2026-04-30",
}
EXCEL_CODE = "\r\n".join(
    (
        "# 受控传输夹具：真实Excel字节，Unicode与CRLF不得被改写。",
        "from io import BytesIO",
        "import requests",
        "from openpyxl import load_workbook",
        "def main(params):",
        f"    response = requests.get({DOWNLOAD!r} + params['fileMd5'], timeout=10)",
        "    response.raise_for_status()",
        "    book = load_workbook(BytesIO(response.content), data_only=True)",
        "    try:",
        "        sheet = book['统计表']",
        "        record = dict(zip(('indicator_name', 'indicator_data',",
        "                           'indicator_unit', 'end_date'),",
        "                          [sheet.cell(2, column).value for column in range(1, 5)]))",
        "    finally:",
        "        book.close()",
        "    return {'code': 0, 'msg': 'success', 'ExcelData': [record]}",
        "",
    )
)
HTML_CODE = "\r\n".join(
    (
        "# 受控HTML链夹具：实际输入为本次markdownMD5对应的Markdown。",
        "import requests",
        "def main(params):",
        f"    response = requests.get({DOWNLOAD!r} + params['markdownMD5'], timeout=10)",
        "    response.raise_for_status()",
        "    row = response.content.decode('utf-8').splitlines()[-1].strip('|').split('|')",
        "    record = dict(zip(('indicator_name', 'indicator_data',",
        "                       'indicator_unit', 'end_date'), [cell.strip() for cell in row]))",
        "    return {'code': 0, 'msg': 'success', 'ExcelData': [record]}",
        "",
    )
)
PROMPT = "\r\n".join(
    (
        "# 来源：V4受控传输夹具；规则来源：测试夹具，不是生产规则。",
        "只读取本次Markdown与GZKJG。保留真实名称及括号单位。",
        "识别表头‘指标、数值、单位、统计期间’，逐条绑定同行字段。",
        "单位来自表格‘单位’列；无单位使用'-'。日期来自‘统计期间’列。",
        "先从真实Markdown构造候选，再以允许的标准化与GZK精确唯一匹配。",
        "不造数据、不改精度、不根据GZK推断单位或日期。",
        "只输出JSON，外层out_data；每条恰有indicator_code、indicator_name、",
        "indicator_data、indicator_unit、end_date五字段。匹配后的代码和名称来自GZK。",
        "本夹具只验证提示词交付；平台453仍负责实际模型提取。",
        "",
    )
)
CASES = (
    ("excel", "xlsx", "fileMd5", "python_excel", EXCEL_CODE),
    ("html", "html", "markdownMD5", "python_html", HTML_CODE),
    ("markdown", "pdf", "markdownMD5", "markdown_prompt", PROMPT),
)


def _serve_v4(settings, sock, stop):
    server = uvicorn.Server(uvicorn.Config(create_app(settings), log_level="error"))

    def stop_later():
        stop.wait()
        server.should_exit = True

    threading.Thread(target=stop_later, daemon=True).start()
    server.run(sockets=[sock])


def _node(filename):
    spec = importlib.util.spec_from_file_location("v4_live_node_" + uuid4().hex, NODES / filename)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _submit(session, base, route, payload):
    response = session.post(base + f"/v4/{route}/jobs", json=payload, timeout=5)
    assert response.status_code == 202, response.text
    return response.json()["job_id"]


def _ready(session, base, job_id):
    deadline = time.monotonic() + 20
    while True:
        response = session.get(base + f"/v4/jobs/{job_id}", params={"wait_seconds": 0.5}, timeout=5)
        assert response.status_code == 200, response.text
        body = response.json()
        assert set(body) == {"agent"} and "artifact_kind" not in body
        item = body["agent"][0]
        assert "artifact_kind" not in item
        if item["status"] != "processing":
            assert item["status"] == "ready", item
            return item
        assert time.monotonic() < deadline, "Controlled spawn worker did not finish"


def _cleanup_owned(directory):
    # Only remove this test's absolute, verified child of runtime/test_runs.
    owned = directory.resolve()
    parent = (ROOT / "runtime" / "test_runs").resolve()
    assert owned.parent == parent and owned.name.startswith("v4-remove-kind-e2e-")
    if owned.exists():
        shutil.rmtree(owned)
    assert not owned.exists()


@pytest.fixture(scope="module")
def v4_live(request):
    # Keep even short-lived test data in the project-owned, identifiable test root.
    directory = ROOT / "runtime" / "test_runs" / ("v4-remove-kind-e2e-" + uuid4().hex)
    request.addfinalizer(lambda: _cleanup_owned(directory))
    settings = Settings(
        _env_file=None,
        runtime_dir=directory / "runtime",
        code_dir=directory / "code",
        api_token=TOKEN,
        model_base_url="",
        model_name="",
        model_api_key="",
        api_worker_poll_interval=0.05,
    )
    repo = Repository(settings.resolved_runtime())
    requests_by_route = {}
    try:
        for route, file_type, source_key, kind, content in CASES:
            payload = {
                source_key: f"v4-controlled-{route}-source",
                "pageTitle": f"V4受控HTTP传输测试\\{route}",
                "fileType": file_type,
                "title": "本次统计标题2026年4月",
                "publictime": "2026-05-01",
                "GZKJG": [
                    {"rule_indicator_name": RECORD["indicator_name"], "rule_indicator_code": 1}
                ],
            }
            task = Task.model_validate(payload)
            report = settings.resolved_runtime() / "fixtures" / route / "quality_report.json"
            report.parent.mkdir(parents=True, exist_ok=True)
            report.write_text(
                json.dumps({"scope": "controlled-v4-transport-only", "credential_probe": TOKEN}),
                encoding="utf-8",
            )
            lease = repo.claim(task.source_key, 60)
            assert lease
            try:
                ArtifactStore(settings, repo).publish(
                    task,
                    lease,
                    content,
                    {
                        "version": f"controlled-v4-{route}-v1",
                        "kind": kind,
                        "contract": (
                            "prompt-out_data-five-fields-v1"
                            if kind == "markdown_prompt"
                            else "main-ExcelData-v1"
                        ),
                        "quality_status": "accepted",
                        "validation_scope": "reviewed-transport-fixture-only",
                        "diagnostics": {
                            "quality_report": report.relative_to(
                                settings.resolved_runtime()
                            ).as_posix()
                        },
                    },
                )
            finally:
                repo.release(task.source_key, lease)
            requests_by_route[route] = payload
    finally:
        repo.close()
    sock = socket.socket()
    sock.bind(("127.0.0.1", 0))
    base = f"http://127.0.0.1:{sock.getsockname()[1]}"
    context = multiprocessing.get_context("spawn")
    stop = context.Event()
    process = context.Process(target=_serve_v4, args=(settings, sock, stop))
    session = requests.Session()
    session.trust_env = False
    session.headers["Authorization"] = "Bearer " + TOKEN
    process.start()
    try:
        deadline = time.monotonic() + 20
        while True:
            try:
                response = session.get(base + "/health", timeout=0.3)
                if response.status_code == 200 and response.json()["status"] == "ok":
                    break
            except requests.RequestException:
                pass
            assert time.monotonic() < deadline, "V4 loopback API/spawn worker startup failed"
            time.sleep(0.05)
        yield base, session, requests_by_route
    finally:
        session.close()
        stop.set()
        process.join(timeout=20)
        if process.is_alive():
            terminate_worker_tree(process)
        sock.close()
    assert process.exitcode == 0


def test_v4_live_health_authentication_and_source_route_boundaries(v4_live):
    base, session, payloads = v4_live
    assert session.get(base + "/v4/health", timeout=3).json() == {
        "status": "ok",
        "protocol_version": "4",
    }
    with requests.Session() as unauthenticated:
        unauthenticated.trust_env = False
        assert unauthenticated.get(base + "/health", timeout=3).status_code == 200
        for route in payloads:
            response = unauthenticated.post(
                base + f"/v4/{route}/jobs", json=payloads[route], timeout=3
            )
            assert response.status_code == 401
        assert unauthenticated.get(base + "/v4/jobs/unknown", timeout=3).status_code == 401
    for route, payload in payloads.items():
        wrong_type = "pdf" if route != "markdown" else "xlsx"
        response = session.post(
            base + f"/v4/{route}/jobs", json={**payload, "fileType": wrong_type}, timeout=3
        )
        assert response.status_code == 422, response.text
        identifier = "fileMd5" if route == "excel" else "markdownMD5"
        missing_identifier = {key: value for key, value in payload.items() if key != identifier}
        assert (
            session.post(base + f"/v4/{route}/jobs", json=missing_identifier, timeout=3).status_code
            == 422
        )
    assert session.get(base + "/v4/jobs/unknown", timeout=3).status_code == 404


@pytest.mark.parametrize(
    "route,file_type,source_key,kind,content", CASES, ids=[row[0] for row in CASES]
)
def test_v4_live_three_route_delivery_and_input_idempotency(
    v4_live, route, file_type, source_key, kind, content
):
    base, session, payloads = v4_live
    payload = payloads[route]
    job_id = _submit(session, base, route, payload)
    item = _ready(session, base, job_id)
    assert tuple(item) == FIELDS
    assert len(item) == 16 and all(isinstance(value, str) for value in item.values())
    assert "artifact_kind" not in item
    assert item["artifact_version"] == f"controlled-v4-{route}-v1"
    assert item["cache_hit"] == "true" and item["quality_status"] == "accepted"
    body_key, other_key = (
        ("PromptCode", "PythonCode") if route == "markdown" else ("PythonCode", "PromptCode")
    )
    assert item[body_key] == content and item[other_key] == ""
    assert item["code_sha256"] == hashlib.sha256(content.encode("utf-8")).hexdigest()
    assert item["report_ref"].startswith(f"/v4/jobs/{job_id}/reports/")
    report = session.get(base + item["report_ref"], timeout=3)
    assert report.status_code == 200 and TOKEN not in report.text
    assert report.json() == {
        "scope": "controlled-v4-transport-only",
        "credential_probe": "[redacted]",
    }
    assert (
        session.get(base + f"/v4/jobs/{job_id}/reports/not-allowed", timeout=3).status_code == 404
    )
    assert _submit(session, base, route, payload) == job_id

    # New collection bytes or a new title are new jobs; the source asset stays reusable.
    for change in ({source_key: payload[source_key] + "-new"}, {"title": "本次统计标题2026年5月"}):
        changed_job = _submit(session, base, route, {**payload, **change})
        assert changed_job != job_id
        changed = _ready(session, base, changed_job)
        assert (
            changed[body_key] == content and changed["artifact_version"] == item["artifact_version"]
        )
        assert changed["cache_hit"] == "true"

    stable = {**payload, "workflow_run_id": f"controlled-v4-{route}-stable-run"}
    stable_job = _submit(session, base, route, stable)
    assert _submit(session, base, route, stable) == stable_job
    _ready(session, base, stable_job)
    response = session.post(
        base + f"/v4/{route}/jobs", json={**stable, "title": "conflicting-title"}, timeout=3
    )
    assert response.status_code == 409


def test_v4_live_v3_excel_compatibility(v4_live):
    base, session, payloads = v4_live
    current = _submit(session, base, "excel", payloads["excel"])
    _ready(session, base, current)
    response = session.post(base + "/v3/excel/jobs", json=payloads["excel"], timeout=3)
    assert response.status_code == 202 and response.json()["job_id"] == current
    legacy = session.get(base + f"/v3/excel/jobs/{current}", timeout=3).json()["agent"][0]
    assert tuple(legacy) == FIELDS[:14]
    assert legacy["PythonCode"] == EXCEL_CODE
    assert legacy["report_ref"].startswith(f"/v3/excel/jobs/{current}/reports/")


def test_v4_live_all_platform_access_nodes_and_python_executors(v4_live, monkeypatch):
    base, _session, payloads = v4_live
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "统计表"
    sheet.append(["指标", "数值", "单位", "统计期间"])
    sheet.append(list(RECORD.values()))
    stream = BytesIO()
    book.save(stream)
    book.close()
    sources = {
        DOWNLOAD + payloads["excel"]["fileMd5"]: stream.getvalue(),
        DOWNLOAD + payloads["html"]["markdownMD5"]: (
            "|指标|数值|单位|统计期间|\n|---|---|---|---|\n|" + "|".join(RECORD.values()) + "|"
        ).encode("utf-8"),
    }
    downloaded = []
    get = requests.get

    def controlled_download(url, *args, **kwargs):
        if url in sources:
            downloaded.append(url)
            response = requests.Response()
            response.status_code = 200
            response._content = sources[url]
            return response
        # API requests stay real; unexpected source downloads never leave the machine.
        assert str(url).startswith(base + "/"), "Unexpected non-loopback download"
        return get(url, *args, **kwargs)

    monkeypatch.setattr(requests, "get", controlled_download)
    access_names = {
        "excel": "01_Excel外部智能体接入_3708.py",
        "html": "02_HTML外部智能体接入_763.py",
        "markdown": "03_Markdown提示词外部智能体接入_2204.py",
    }
    executor_names = {
        "excel": "04_Excel平台执行器_4912.py",
        "html": "05_HTML平台执行器_5242.py",
    }
    for route, payload in payloads.items():
        access = _node(access_names[route])
        access.API_BASE_URL, access.API_TOKEN = base, TOKEN
        access.POLL_SECONDS = 0.1
        # Real service rejects extra fields: the copyable node must omit the
        # unused source ID, including the reported xlsx + markdownMD5=null case.
        access_params = {
            **payload,
            **({"markdownMD5": None} if route == "excel" else {"fileMd5": "unused-raw-id"}),
        }
        delivered = access.main(access_params)
        assert "artifact_kind" not in delivered
        item = delivered["agent"][0]
        assert item["status"] == "ready"
        assert tuple(item) == tuple(field for field in FIELDS if field != "report_ref")
        assert len(item) == 15 and "report_ref" not in item and "artifact_kind" not in item
        assert all(isinstance(value, str) for value in item.values())
        assert access.main(access_params)["agent"][0]["job_id"] == item["job_id"]
        if route == "markdown":
            assert delivered["PromptFile"] == item["PromptFile"]
            assert delivered["PromptCode"] == PROMPT == item["PromptCode"]
            assert item["PythonCode"] == item["PythonFile"] == ""
        else:
            business = {key: value for key, value in payload.items() if key != "GZKJG"}
            executor = _node(executor_names[route])
            assert executor.main({**business, **delivered}) == {
                "code": 0,
                "msg": "success",
                "ExcelData": [RECORD],
            }
    assert downloaded == list(sources)
