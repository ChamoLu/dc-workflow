"""End-to-end reference propagation and nonblocking completeness with real execution."""

import json
from pathlib import Path

import openpyxl
import pytest
from test_application import ScriptedModel, make_service
from test_application import inputs as inputs

from dc_agent.domain import Design, digest
from dc_agent.validation import check_design


@pytest.mark.parametrize("duplicate", [False, True])
def test_catalog_gate_precedes_source_and_all_model_calls(inputs, duplicate):
    settings, task = inputs
    path = settings.resolved_rules() / "todo.xlsx"
    book = openpyxl.load_workbook(path)
    if duplicate:
        book.active.append([c.value for c in book.active[2]])
    else:
        book.active["E2"] = "different-title"
    book.save(path)
    book.close()
    model = ScriptedModel()
    service = make_service(settings, model)

    class NoSource:
        def read(self, task):
            pytest.fail("Catalog failure must stop before source download")

    service.source = NoSource()
    try:
        result = service.run(task)
        assert result.status == "failed", result
        assert result.error["code"] == (
            "pageTitle_redundancy" if duplicate else "todoFile_no_pageTitle"
        )
        assert result.error["stage"] == "resolve_rule_context"
        assert model.calls == [] and model.preflight_calls == 0
        assert service.repo.current(task.source_key) is None
    finally:
        service.close()


def test_same_task_and_reference_snapshot_reaches_design_repairs_and_review(
    inputs,
    write_reference,
):
    settings, task = inputs
    write_reference(
        settings.resolved_rules() / "IndicatorName_Excel_Md.xlsx",
        "reference-code",
        rows=[("参考可能错误", "999"), ("说明", "以实际源格位为准")],
    )

    class InspectingModel(ScriptedModel):
        def __init__(self):
            super().__init__(wrong_first=True)
            self.contexts = []

        def complete(self, operation, system, context, *args):
            if operation == "preflight":
                assert "task_info" not in context and "references" not in context
            else:
                self.contexts.append((context["task_info"], context["references"]))
                assert context["asset_contract"]["asset_kind"] == "python_excel"
                assert "Do not filter source records" in context["asset_contract"]["gzk_usage"]
                assert "本次产物契约" in system
                assert context["task_info"]["fields"]["数据目标表"] == "test_table"
                assert context["references"][0]["cells"][0]["value"] == "参考可能错误"
            result = super().complete(operation, system, context, *args)
            if operation == "design" or operation.startswith("review"):
                result["reference_conflicts"] = ["参考B1写999，实际Sheet!B2为0；采用真实源数据0"]
            return result

    model = InspectingModel()
    service = make_service(settings, model)
    try:
        result = service.run(task)
        assert result.status == "ready", result
        assert model.calls == ["design", "generate_0", "generate_1", "review_1"]
        assert all(pair == model.contexts[0] for pair in model.contexts)
        metadata = result.artifact
        snapshot = json.loads(
            (settings.resolved_runtime() / metadata["generation_context_path"]).read_text(
                encoding="utf-8"
            )
        )
        assert digest(snapshot) == metadata["generation_context_sha256"]
        report = json.loads(
            (settings.resolved_runtime() / metadata["validation_path"]).read_text(encoding="utf-8")
        )
        assert report["records"][0]["indicator_data"] == "0"
        assert any(c["name"] == "reference_conflicts" for c in report["checks"])
    finally:
        service.close()


def test_reference_change_rebuilds_saved_design_and_rechecks_saved_candidate(
    inputs, write_reference
):
    settings, task = inputs
    model = ScriptedModel()
    service = make_service(settings, model)
    try:
        first = service.run(task)
        assert first.status == "ready", first
        write_reference(
            settings.resolved_rules() / "IndicatorName_Excel_Md.xlsx",
            "reference-code",
            rows=[("维护后新增的首行参考", "规则")],
        )
        model.calls.clear()
        second = service.run(task, retry=True, rebuild=True, resume_candidate=True)
        assert second.status == "ready", second
        assert model.calls == ["design", "review_0"]
        assert (
            second.artifact["generation_context_sha256"]
            != first.artifact["generation_context_sha256"]
        )
        assert second.artifact["version"] != first.artifact["version"]
    finally:
        service.close()


def test_cache_reuse_does_not_read_current_catalog_but_rebuild_does(inputs):
    settings, task = inputs
    model = ScriptedModel()
    service = make_service(settings, model)
    try:
        first = service.run(task)
        assert first.status == "ready", first
        (settings.resolved_rules() / "todo.xlsx").unlink()
        next_task = task.model_copy(update={"request_id": "cache-after-rule-edit"})
        cached = service.run(next_task)
        assert cached.status == "ready" and cached.cache_hit
        assert model.preflight_calls == 1 and len(model.calls) == 3
        rebuilt = service.run(next_task, retry=True, rebuild=True)
        assert rebuilt.error["code"] == "REFERENCE_UNAVAILABLE"
        assert service.repo.current(task.source_key)["version"] == first.artifact["version"]
    finally:
        service.close()


@pytest.mark.parametrize(
    "verdict",
    [
        "approved",
        "coverage_only",
        "correctness",
        "invalid_refs",
        "reference_citation",
        "bad_reference",
    ],
)
def test_partial_design_and_review_issues_publish_with_explicit_quality_status(
    inputs,
    write_reference,
    verdict,
):
    settings, task = inputs
    settings = settings.model_copy(update={"max_repairs": 0})
    path = Path(task.source_path)
    book = openpyxl.load_workbook(path)
    book.create_sheet("另一个有效表").append(["另一指标", 42])
    book.save(path)
    book.close()
    write_reference(
        settings.resolved_rules() / "IndicatorName_Excel_Md.xlsx",
        "reference-code",
        rows=[("参考片段", "参考名称")],
    )

    class PartialModel(ScriptedModel):
        def complete(self, operation, *args):
            reply = super().complete(operation, *args)
            if operation.startswith("review"):
                reply["quality_findings"] = ["另一个有效表有数据尚未提取"]
                if verdict in {"coverage_only", "invalid_refs"}:
                    reply.update(approved=False, rejection_kind="coverage", findings=["提取不全"])
                if verdict == "correctness":
                    reply.update(
                        approved=False, rejection_kind="correctness", findings=["输出名称归属错误"]
                    )
                if verdict == "invalid_refs":
                    reply["reviewed_refs"] = ["fabricated-source!Z99"]
                if verdict == "reference_citation":
                    reply["reviewed_refs"].append(
                        "IndicatorName_Excel_Md.xlsx:reference-code!A1:B1"
                    )
                if verdict == "bad_reference":
                    reply["reviewed_refs"].append("IndicatorName_Excel_Md.xlsx:reference-code!Z99")
            return reply

    model = PartialModel()
    service = make_service(settings, model)
    try:
        result = service.run(task)
        assert result.status == "ready", result
        if verdict in {"correctness", "invalid_refs", "bad_reference"}:
            assert result.artifact["quality_status"] == "usable_with_warnings"
            assert result.artifact["selection"]["semantic_status"] != "accepted"
            assert Path(result.diagnostics["quality_report"]).is_file()
        else:
            assert result.status == "ready", result
            report = json.loads(
                (settings.resolved_runtime() / result.artifact["validation_path"]).read_text(
                    encoding="utf-8"
                )
            )
            assert len(report["records"]) == 1
            warnings = {c["name"] for c in report["checks"] if c["status"] == "warning"}
            assert {
                "design_section_coverage",
                "review_section_coverage",
                "extraction_quality",
            } <= warnings
        assert model.calls == ["design", "generate_0", "review_0"]
    finally:
        service.close()


def test_design_partial_coverage_is_advisory_but_unknown_source_is_invalid():
    design = Design(
        summary="partial",
        name_logic="source labels",
        unit_logic="source units",
        unit_position="A1",
        date_logic="unknown",
        date_position="unknown",
        covered_sections=[],
    )
    inspection = {"sections": [{"id": "real-sheet"}]}
    assert check_design(design, inspection)[0].status == "warning"
    with pytest.raises(ValueError, match="unknown source"):
        check_design(design.model_copy(update={"covered_sections": ["invented-sheet"]}), inspection)
