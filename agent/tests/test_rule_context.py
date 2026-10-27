"""Catalog identity, maintained-file freshness and optional worksheet routing."""

import os

import openpyxl
import pytest

from dc_agent.config import Settings
from dc_agent.domain import AgentError, Task
from dc_agent.references import WorkbookRuleContext
from dc_agent.references.evidence import reference_ref_exists


@pytest.fixture
def rule_settings(tmp_path):
    return Settings(_env_file=None, rule_dir=tmp_path / "rule")


def task_for(file_type="xlsx", title="来源\\统计\\月报"):
    return Task(
        pageTitle=title,
        fileType=file_type,
        fileMd5="file-id",
        markdownMD5="markdown-id",
        GZKJG=[{"rule_indicator_name": "源指标"}],
    )


def catalog_row(task, code="0012345678901234567"):
    return [
        "000012345",
        "测试来源",
        code,
        "示例不是指标全集",
        task.pageTitle,
        "依据本次真实表头",
        "全角转半角",
        task.fileType,
        "模块",
        "target_table",
    ]


@pytest.mark.parametrize(
    "file_type",
    [
        "xls",
        "xlsx",
        "csv",
        "html",
        "htm",
        "shtml",
        "md",
        "markdown",
        "pdf",
        "png",
        "jpg",
        "jpeg",
        "doc",
        "docx",
    ],
)
def test_routes_types_to_exact_sheet_without_skipping_first_row(
    rule_settings,
    write_catalog,
    write_reference,
    file_type,
):
    task = task_for(file_type)
    row = catalog_row(task)
    directory = rule_settings.resolved_rules()
    write_catalog(directory, [row], extra_headers=["未来新增字段"])
    for name, marker in [
        ("IndicatorName_Excel_Md.xlsx", "名称片段"),
        ("rule_html.xlsx", "正则参考"),
    ]:
        write_reference(
            directory / name,
            row[2],
            rows=[(marker, "首行内容"), ("另一个指标", "第二行")],
        )
    result = WorkbookRuleContext(rule_settings).load(task)
    assert result["task_info"]["fields"]["数据来源编码"] == "000012345"
    assert result["task_info"]["fields"]["数据目标表"] == "target_table"
    assert result["task_info"]["data_source_path_code"] == row[2]
    (reference,) = result["references"]
    assert reference["sheet"] == row[2]
    assert reference["cells"][0] == {
        "address": "A1",
        "value": "正则参考" if task.kind == "python_html" else "名称片段",
    }
    assert reference["cells"][-1] == {"address": "B2", "value": "第二行"}


@pytest.mark.parametrize(
    "different_title",
    [
        "来源_统计_月报",
        "来源/统计/月报",
        "来源\\统计\\月报 ",
        "来源\\统计\\月",
    ],
)
def test_page_title_is_exact_and_does_not_use_filename_normalization(
    rule_settings,
    write_catalog,
    different_title,
):
    task = task_for()
    row = catalog_row(task)
    row[4] = different_title
    write_catalog(rule_settings.resolved_rules(), [row])
    with pytest.raises(AgentError) as error:
        WorkbookRuleContext(rule_settings).load(task)
    assert error.value.code == "todoFile_no_pageTitle"


@pytest.mark.parametrize("across_sheets", [False, True])
def test_duplicate_title_is_rejected_even_for_different_types_and_codes(
    rule_settings,
    write_catalog,
    across_sheets,
):
    task = task_for()
    first = catalog_row(task)
    second = catalog_row(task_for("html"), code="different-code")
    path = write_catalog(
        rule_settings.resolved_rules(), [first, *([] if across_sheets else [second])]
    )
    if across_sheets:
        book = openpyxl.load_workbook(path)
        sheet = book.copy_worksheet(book.active)
        sheet.title = "新增任务表"
        for column, value in enumerate(second, 1):
            sheet.cell(2, column, value)
        book.save(path)
        book.close()
    with pytest.raises(AgentError) as error:
        WorkbookRuleContext(rule_settings).load(task)
    assert error.value.code == "pageTitle_redundancy"
    assert len(error.value.details["matching_rows"]) == 2


@pytest.mark.parametrize("code", ["not-present", None, "", 1234567890123456789])
def test_missing_or_unusable_sheet_identifier_proceeds_with_task_information_only(
    rule_settings,
    write_catalog,
    write_reference,
    code,
):
    task = task_for()
    directory = rule_settings.resolved_rules()
    write_catalog(directory, [catalog_row(task, code=code)])
    write_reference(directory / "IndicatorName_Excel_Md.xlsx", "unrelated")
    result = WorkbookRuleContext(rule_settings).load(task)
    assert set(result) == {"task_info", "references"}
    assert result["references"] == []
    assert result["task_info"]["fields"]["pageTitle"] == task.pageTitle


def test_new_tasks_and_changed_references_are_seen_by_same_provider(
    rule_settings,
    write_catalog,
    write_reference,
):
    directory = rule_settings.resolved_rules()
    first, second = task_for(), task_for(title="新增来源\\新增任务")
    catalog = write_catalog(directory, [catalog_row(first)])
    reference = write_reference(directory / "IndicatorName_Excel_Md.xlsx", catalog_row(first)[2])
    provider = WorkbookRuleContext(rule_settings)
    old = provider.load(first)
    previous_stat = reference.stat()
    write_reference(reference, catalog_row(first)[2], rows=[("维护后内容", "新片段")])
    # Even timestamps restored by sync/copy software must not cause stale data.
    os.utime(reference, ns=(previous_stat.st_atime_ns, previous_stat.st_mtime_ns))
    fresh = provider.load(first)
    assert fresh["references"][0]["cells"][0]["value"] == "维护后内容"
    assert fresh["references"][0]["sha256"] != old["references"][0]["sha256"]
    write_catalog(directory, [catalog_row(first), catalog_row(second, code="new-code")])
    added = provider.load(second)
    assert catalog.is_file()
    assert added["task_info"]["source"]["row"] == 3
    assert added["references"] == []


def test_configurable_filenames_and_task_sheet_name(rule_settings, write_catalog, write_reference):
    task = task_for("html")
    directory = rule_settings.resolved_rules()
    path = write_catalog(directory, [catalog_row(task)], sheet_name="随时可改的任务Sheet")
    path.rename(directory / "tasks-current.xlsx")
    write_reference(directory / "html-current.xlsx", catalog_row(task)[2])
    settings = rule_settings.model_copy(
        update={
            "todo_filename": "tasks-current.xlsx",
            "html_reference_filename": "html-current.xlsx",
        }
    )
    result = WorkbookRuleContext(settings).load(task)
    assert result["task_info"]["source"]["sheet"] == "随时可改的任务Sheet"
    assert result["references"][0]["file"] == "html-current.xlsx"


@pytest.mark.parametrize(
    "problem,code",
    [
        ("missing_catalog", "REFERENCE_UNAVAILABLE"),
        ("bad_catalog", "REFERENCE_FORMAT"),
        ("wrong_headers", "REFERENCE_FORMAT"),
        ("missing_reference", "REFERENCE_UNAVAILABLE"),
        ("bad_reference", "REFERENCE_FORMAT"),
    ],
)
def test_file_or_schema_failures_are_distinct_from_absent_task_or_sheet(
    rule_settings,
    write_catalog,
    problem,
    code,
):
    task = task_for()
    directory = rule_settings.resolved_rules()
    path = write_catalog(directory, [catalog_row(task)])
    if problem == "missing_catalog":
        path.unlink()
    elif problem == "bad_catalog":
        path.write_bytes(b"not an xlsx")
    elif problem == "wrong_headers":
        book = openpyxl.load_workbook(path)
        book.active["E1"] = "wrong"
        book.save(path)
        book.close()
    elif problem == "bad_reference":
        (directory / "IndicatorName_Excel_Md.xlsx").write_bytes(b"not an xlsx")
    with pytest.raises(AgentError) as error:
        WorkbookRuleContext(rule_settings).load(task)
    assert error.value.code == code
    assert error.value.stage == "resolve_rule_context"


@pytest.mark.parametrize(
    "citation,expected",
    [
        ("names.xlsx:00123!A1", True),
        ("names.xlsx:00123!A1:K36", True),
        ("names.xlsx:00123!B1", False),
        ("names.xlsx:00123!A1:Z99", False),
        ("names.xlsx:00123!A:K", False),
        ("other.xlsx:00123!A1", False),
        ("names.xlsx:0123!A1", False),
        ("fabricated-source!A1", False),
    ],
)
def test_reference_review_citations_require_selected_file_sheet_and_bounded_content(
    citation, expected
):
    references = [
        {
            "file": "names.xlsx",
            "sheet": "00123",
            "cells": [{"address": "A1", "value": "片段"}, {"address": "K36", "value": "名称"}],
        }
    ]
    assert reference_ref_exists(citation, references) is expected
