"""Use real in-memory containers; do not substitute converted XLS fixtures."""

import gzip
import io
import json
import zipfile
from datetime import datetime

import openpyxl
import pytest

from dc_agent.inspection import InspectionError, inspect_bytes


def workbook_bytes():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "统计"
    sheet["A1"] = "指标"
    sheet["B1"] = 2010
    sheet["A2"] = "生产总值(亿元)"
    sheet["B2"] = 0
    sheet["C2"] = datetime(2026, 10, 15)
    sheet.merge_cells("A3:C3")
    sheet["A3"] = "单位:亿元"
    workbook.create_sheet("空表")
    target = io.BytesIO()
    workbook.save(target)
    return target.getvalue()


def test_all_sheets_merges_zero_and_date():
    report = inspect_bytes(workbook_bytes(), "excel")
    assert report["format"] == "xlsx"
    assert report["coverage_complete"] is True
    assert len(report["sections"]) == 2
    sheet, empty = report["sections"]
    assert sheet["merged_ranges"] == ["A3:C3"]
    cells = {cell["ref"]: cell for cell in sheet["cells"]}
    assert cells["统计!B2"]["value"] == 0
    assert cells["统计!C2"]["python_type"] == "datetime"
    assert cells["统计!C2"]["value"] == "2026-10-15T00:00:00"
    assert empty["kind"] == "empty"
    json.dumps(report, allow_nan=False)


def test_gzip_magic_and_grid_limit():
    data = gzip.compress(workbook_bytes())
    assert inspect_bytes(data, "xls")["format"] == "xlsx"
    with pytest.raises(InspectionError, match="max_cells"):
        inspect_bytes(data, "xlsx", max_cells=3)


def test_formula_without_cache_is_preserved_as_unknown_value():
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "=1+1"
    target = io.BytesIO()
    workbook.save(target)
    report = inspect_bytes(target.getvalue(), "xlsx")
    cell = report["sections"][0]["cells"][0]
    assert cell["value"] is None
    assert cell["formula"] == "=1+1"
    assert cell["cache_missing"] is True
    assert report["sections"][0]["kind"] == "table"


def test_csv_quotes_zero_blank_and_encoding():
    content = '项目,数值,说明\r\n甲,0,"a,b"\r\n乙,,空值\r\n'
    report = inspect_bytes(content.encode("gb18030"), "csv")
    cells = {cell["ref"]: cell["value"] for cell in report["sections"][0]["cells"]}
    assert cells["csv!B2"] == "0"
    assert cells["csv!C2"] == "a,b"
    assert "csv!B3" not in cells
    with pytest.raises(InspectionError, match="max_cells"):
        inspect_bytes(content.encode(), "csv", max_cells=3)


def test_markdown_preserves_headers_notes_and_empty_lines():
    content = "# 月报\n\n| 指标 | 1-5月 |\n|---|---|\n| 产值 | 0 |\n单位:万元\n"
    report = inspect_bytes(content.encode(), "md")
    assert report["sections"][0]["cells"][1]["value"] == ""
    assert report["sections"][0]["cells"][-1] == {
        "ref": "L6",
        "value": "单位:万元",
        "python_type": "str",
    }
    with pytest.raises(InspectionError, match="max_cells"):
        inspect_bytes(content.encode(), "md", max_cells=2)


def test_invalid_and_unknown_formats():
    for data, file_type in [
        (b"PKbad", "xlsx"),
        (b"\x1f\x8bbad", "md"),
        (b"\x00\x01", "csv"),
        (b"{}", "json"),
    ]:
        with pytest.raises(InspectionError):
            inspect_bytes(data, file_type)


def test_zip_bomb_rejected_before_workbook_loading():
    target = io.BytesIO()
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/workbook.xml", b"x" * 2_000_000)
        archive.writestr("[Content_Types].xml", "x")
    with pytest.raises(InspectionError, match="compression ratio"):
        inspect_bytes(target.getvalue(), "xlsx")


def test_gzip_decompression_has_byte_limit(monkeypatch):
    from dc_agent.inspection import readers as inspection

    monkeypatch.setattr(inspection, "MAX_BYTES", 100)
    with pytest.raises(InspectionError, match="Decompressed"):
        inspect_bytes(gzip.compress(b"x" * 1000), "md")
