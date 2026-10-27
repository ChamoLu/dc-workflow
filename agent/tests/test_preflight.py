import io
import json
import time

import httpx
import openpyxl
import pytest

from dc_agent.application.preflight import ContentPreflight, inspect_source
from dc_agent.config import Settings
from dc_agent.domain import AgentError, Task
from dc_agent.infrastructure import ChatCompletionsGateway, FileSourceProvider


def task_for(kind="xlsx"):
    return Task(
        pageTitle="site",
        fileType=kind,
        fileMd5="file",
        markdownMD5="md",
        GZKJG=[{"rule_indicator_name": "Absent rule name"}],
    )


def real_inspection():
    book = openpyxl.Workbook()
    book.active.title = "目录"
    book.active.append(["统计表", "2026年"])
    sheet = book.create_sheet("数据")
    sheet.append(["指标", "值"])
    sheet.append(["业务甲", 0])
    sheet.append(["业务乙", -1.25])
    buf = io.BytesIO()
    book.save(buf)
    return inspect_source(task_for(), buf.getvalue())


class ContentModel:
    def __init__(self, reply):
        self.reply, self.calls = reply, []

    def complete(self, operation, system, context, schema, timeout):
        assert operation == "preflight"
        assert "GZKJG" not in context and "task" not in context
        self.calls.append(context)
        if isinstance(self.reply, Exception):
            raise self.reply
        return self.reply


@pytest.mark.parametrize("verdict", ["has_data", "uncertain", "no_data"])
def test_coarse_model_verdict_and_retry_journal(tmp_path, verdict):
    settings = Settings(_env_file=None)
    model = ContentModel({"verdict": verdict, "reason": "Coarse source decision"})
    preflight = ContentPreflight(model, settings)
    inspection = real_inspection()
    for _ in range(2):
        if verdict == "no_data":
            with pytest.raises(AgentError) as error:
                preflight.check(task_for(), inspection, tmp_path, time.monotonic() + 30)
            assert error.value.code == "FilenoData"
        else:
            result = preflight.check(task_for(), inspection, tmp_path, time.monotonic() + 30)
            assert result["allowed"] is True
            assert result["verdict"] == verdict
    assert len(model.calls) == 1
    sections = model.calls[0]["inspection"]["sections"]
    assert [s["id"] for s in sections] == ["目录", "数据"]
    assert {c["value"] for c in sections[1]["cells"]} >= {0, -1.25}


@pytest.mark.parametrize(
    "reply,code",
    [
        ({"verdict": "maybe", "reason": "x"}, "MODEL_SCHEMA_INVALID"),
        ({"verdict": False, "reason": "x"}, "MODEL_SCHEMA_INVALID"),
        ({"verdict": "no_data", "reason": " "}, "MODEL_SCHEMA_INVALID"),
        (AgentError("MODEL_TIMEOUT", "model timeout", retryable=True), "MODEL_TIMEOUT"),
    ],
)
def test_invalid_or_failed_model_decision_is_not_no_data(tmp_path, reply, code):
    with pytest.raises(AgentError) as error:
        ContentPreflight(ContentModel(reply), Settings(_env_file=None)).check(
            task_for(), real_inspection(), tmp_path, time.monotonic() + 30
        )
    assert error.value.code == code
    assert not (tmp_path / "preflight.json").exists()


def test_source_change_invalidates_only_preflight_decision(tmp_path):
    model = ContentModel({"verdict": "has_data", "reason": "Source has data"})
    checker = ContentPreflight(model, Settings(_env_file=None))
    original = real_inspection()
    checker.check(task_for(), original, tmp_path, time.monotonic() + 30)
    changed = {**original, "sha256": "different-content-hash"}
    checker.check(task_for(), changed, tmp_path, time.monotonic() + 30)
    assert len(model.calls) == 2


def test_empty_download_reaches_model_but_corrupt_file_does_not(tmp_path):
    source = tmp_path / "empty.xlsx"
    source.write_bytes(b"")
    settings = Settings(_env_file=None)
    task = task_for().model_copy(update={"source_path": str(source)})
    raw = FileSourceProvider(settings).read(task)
    inspection = inspect_source(task, raw)
    assert inspection["coverage_complete"] and inspection["size"] == 0
    model = ContentModel({"verdict": "no_data", "reason": "Empty bytes"})
    with pytest.raises(AgentError, match="有效"):
        ContentPreflight(model, settings).check(task, inspection, tmp_path, time.monotonic() + 30)
    assert len(model.calls) == 1
    with pytest.raises(AgentError) as error:
        inspect_source(task, b"PKbroken")
    assert error.value.code == "FORMAT_UNREADABLE"


def test_markdown_route_uses_actual_text_and_rejects_workbook_bytes():
    task = task_for("html")
    inspection = inspect_source(task, "实际统计正文：业务指标为0。".encode())
    assert inspection["format"] == "markdown"
    book = openpyxl.Workbook()
    stream = io.BytesIO()
    book.save(stream)
    with pytest.raises(AgentError) as error:
        inspect_source(task, stream.getvalue())
    assert error.value.code == "FORMAT_UNREADABLE"


def test_real_gateway_request_uses_preflight_budget(tmp_path):
    settings = Settings(
        _env_file=None,
        model_base_url="https://model.example/v1",
        model_name="test",
        model_api_key="fake",
        model_stream=False,
        model_preflight_max_tokens=321,
    )
    seen = []

    def handler(request):
        body = json.loads(request.content)
        seen.append(body)
        assert body["max_tokens"] == 321
        assert json.loads(body["messages"][1]["content"])["operation"] == "preflight"
        return httpx.Response(
            200,
            json={
                "choices": [
                    {
                        "finish_reason": "stop",
                        "message": {
                            "content": json.dumps(
                                {"verdict": "uncertain", "reason": "Possible data"}
                            )
                        },
                    }
                ]
            },
        )

    gateway = ChatCompletionsGateway(settings, httpx.MockTransport(handler))
    result = ContentPreflight(gateway, settings).check(
        task_for(), real_inspection(), tmp_path, time.monotonic() + 30
    )
    assert result["allowed"] and len(seen) == 1
