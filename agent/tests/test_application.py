import io
import json
from pathlib import Path

import httpx
import openpyxl
import pytest

from dc_agent.application.service import AssetService
from dc_agent.config import Settings
from dc_agent.domain import AgentError, Task, normalize_name
from dc_agent.entrypoints.cli import load_task
from dc_agent.generation import SkillRegistry
from dc_agent.infrastructure import ChatCompletionsGateway, FileSourceProvider, LocalReferences
from dc_agent.infrastructure.execution import LocalProcessRunner
from dc_agent.infrastructure.storage import Repository

CODE = """import requests
import openpyxl
from io import BytesIO

def main(params):
    data = requests.get("http://service-1.example.invalid/api/file/download/" + params["fileMd5"]).content
    sheet = openpyxl.load_workbook(BytesIO(data), data_only=True).worksheets[0]
    return {"code": 0, "msg": "success", "ExcelData": [
        {"indicator_name": sheet["A2"].value, "indicator_data": str(sheet["B2"].value),
         "indicator_unit": "-", "end_date": ""}
    ]}
"""


class ScriptedModel:
    """Hand-written test replies; never advertised as real model generation."""

    def __init__(self, wrong_first=False):
        self.calls = []
        self.preflight_calls = 0
        self.wrong_first = wrong_first

    def complete(self, operation, system, context, schema, timeout):
        if operation == "preflight":
            self.preflight_calls += 1
            assert "GZKJG" not in context
            return {"verdict": "has_data", "reason": "Test source contains a business value"}
        self.calls.append(operation)
        assert "source_rules_v2" in system
        assert "excel_data" not in context.get("task", {})
        if operation == "design":
            return {
                "summary": "Simple two-column real workbook",
                "name_logic": "A2",
                "unit_logic": "No explicit unit; -",
                "unit_position": "not provided",
                "date_logic": "Unknown; blank",
                "date_position": "not provided",
                "covered_sections": ["Sheet"],
                "limitations": ["single-sample"],
            }
        if operation.startswith("generate"):
            assert context["parser_contract"]["download_base_url"].endswith("/api/file/download/")
            assert context["parser_contract"]["download_parameter"] == "fileMd5"
            assert "fileServiceUrl" not in context["parser_contract"]["available_parameters"]
            assert "Finite scalar" in context["parser_contract"]["record_fields"]["indicator_data"]
            assert "Numeric strings" in context["parser_contract"]["unit_binding_policy"]
            return {
                "content": CODE
                if not (self.wrong_first and operation == "generate_0")
                else CODE.replace('"ExcelData"', '"OUTDATA"'),
                "evidence": [
                    {
                        "record_index": 0,
                        "name_refs": ["Sheet!A2"],
                        "data_ref": "Sheet!B2",
                        "unit_refs": [],
                        "date_refs": [],
                    }
                ],
            }
        assert context["review_contract"]["asset_kind"] == "python_excel"
        assert "never require indicator_code" in context["review_contract"]["gzk_policy"]
        assert "numeric strings are both valid" in context["review_contract"]["numeric_policy"]
        assert context["runtime_report"]["records"]
        return {"approved": True, "findings": [], "reviewed_refs": ["Sheet!A2", "Sheet!B2"]}


@pytest.fixture
def inputs(tmp_path, write_catalog, write_reference):
    book = openpyxl.Workbook()
    book.active.append(["指标", "值"])
    book.active.append(["指标甲", 0])
    data = io.BytesIO()
    book.save(data)
    source = tmp_path / "source.xlsx"
    source.write_bytes(data.getvalue())
    skill_root = tmp_path / "skills"
    skill = skill_root / "excel-rule-parser-codegen" / "SKILL.md"
    skill.parent.mkdir(parents=True)
    skill.write_text(
        "Original skill requiring unavailable workbook; adaptation supersedes this.",
        encoding="utf-8",
    )
    settings = Settings(
        _env_file=None,
        runtime_dir=tmp_path / "runtime",
        rule_dir=tmp_path / "rule",
        skill_root=skill_root,
        runner="local",
        allow_unsafe_local_execution=True,
        max_repairs=1,
    )
    task = Task(
        request_id="one",
        pageTitle="站点\\统计",
        fileType="xlsx",
        fileMd5="test-id",
        source_path=str(source),
        GZKJG=[{"rule_indicator_name": "指标甲", "rule_indicator_code": 123}],
    )
    write_catalog(
        settings.resolved_rules(),
        [
            [
                "test-source",
                "测试来源",
                "reference-code",
                "指标甲",
                task.pageTitle,
                "依据真实指标行",
                "全角转半角",
                task.fileType,
                "测试模块",
                "test_table",
            ]
        ],
    )
    write_reference(settings.resolved_rules() / "IndicatorName_Excel_Md.xlsx", code="other-code")
    return settings, task


def make_service(settings, model):
    return AssetService(
        settings,
        model,
        FileSourceProvider(settings),
        LocalReferences(),
        LocalProcessRunner(allow_unsafe=True),
        SkillRegistry(settings.skill_root),
    )


def test_real_bytes_generation_validation_reuse(inputs):
    settings, task = inputs
    model = ScriptedModel()
    service = make_service(settings, model)
    try:
        result = service.run(task)
        assert result.status == "ready", result
        assert model.calls == ["design", "generate_0", "review_0"]
        assert model.preflight_calls == 1
        path = settings.resolved_runtime() / result.artifact["validation_path"]
        report = json.loads(path.read_text(encoding="utf-8"))
        assert report["records"][0]["indicator_data"] == "0"
        assert report["matching"]["matched_records"] == 1
        assert service.run(task).task_id == result.task_id
        another = task.model_copy(update={"request_id": "two"})
        assert service.run(another).cache_hit
        assert len(model.calls) == 3
        assert model.preflight_calls == 1
        assert (
            settings.resolved_runtime() / "tasks" / result.task_id / "checkpoints.sqlite"
        ).exists()
    finally:
        service.close()


@pytest.mark.parametrize("failure", [None, "different_result", "unknown_ref", "missing_helper"])
def test_runtime_evidence_uses_real_production_output_and_checks_refs(inputs, failure):
    settings, task = inputs
    settings = settings.model_copy(update={"max_repairs": 0})

    class RuntimeModel(ScriptedModel):
        def complete(self, *args):
            reply = super().complete(*args)
            if args[0].startswith("generate"):
                helper = """
def _dc_validate(params):
    result = main(params)
    result["_dc_evidence"] = [{"record_index": 0, "name_refs": ["Sheet!A2"], "data_ref": "Sheet!B2", "unit_refs": [], "date_refs": []}]
    return result
"""
                if failure == "different_result":
                    helper = helper.replace(
                        "    return result",
                        '    result["ExcelData"][0]["indicator_data"] = "999"\n    return result',
                    )
                if failure == "unknown_ref":
                    helper = helper.replace("Sheet!B2", "Sheet!Z99")
                reply.update(
                    content=CODE + (helper if failure != "missing_helper" else ""),
                    evidence_mode="runtime",
                    evidence=[],
                )
                assert args[2]["parser_contract"]["evidence_protocol"]["mode"] == "runtime"
            return reply

    service = make_service(settings, RuntimeModel())
    try:
        result = service.run(task)
        if failure is None:
            assert result.status == "ready", result
            report = json.loads(
                (settings.resolved_runtime() / result.artifact["validation_path"]).read_text(
                    encoding="utf-8"
                )
            )
            assert report["evidence"][0]["data_ref"] == "Sheet!B2"
            assert report["records"][0]["indicator_data"] == "0"
            assert "_dc_evidence" not in report["records"][0]
        else:
            assert result.status == "review_required", result
            assert service.repo.current(task.source_key) is None
    finally:
        service.close()


def test_repair_and_idempotency_conflict(inputs):
    settings, task = inputs
    model = ScriptedModel(wrong_first=True)
    service = make_service(settings, model)
    try:
        assert service.run(task).status == "ready"
        assert model.calls == ["design", "generate_0", "generate_1", "review_1"]
        with pytest.raises(AgentError, match="different input"):
            service.run(task.model_copy(update={"title": "different"}))
    finally:
        service.close()


def test_resume_latest_candidate_revalidates_and_has_new_repair_budget(inputs):
    settings, task = inputs
    first = make_service(settings.model_copy(update={"max_repairs": 0}), ScriptedModel(True))
    try:
        failed = first.run(task)
        assert failed.status == "review_required"
    finally:
        first.close()
    model = ScriptedModel()
    service = make_service(settings, model)
    try:
        ready = service.run(task, retry=True, resume_candidate=True)
        assert ready.status == "ready", ready
        assert model.calls == ["generate_1", "review_1"]
        workspace = settings.resolved_runtime() / "tasks" / ready.task_id
        assert (workspace / "candidate-0.json").exists()
        assert (workspace / "candidate-1.json").exists()
        service.repo.deactivate(task.source_key)
        model.calls.clear()
        result = service.run(task, retry=True, resume_candidate=True)
        assert result.status == "ready", result
        # Latest saved candidate passes; no design or generation replay is needed.
        assert not any(call.startswith(("design", "generate")) for call in model.calls)
    finally:
        service.close()


def test_resume_rejects_changed_source_and_missing_candidate(inputs):
    settings, task = inputs
    model = ScriptedModel()
    service = make_service(settings, model)
    try:
        missing = service.run(task, retry=True, resume_candidate=True)
        assert missing.error["code"] == "RESUME_UNAVAILABLE"
        ready = service.run(task, retry=True)
        assert ready.status == "ready"
        service.repo.deactivate(task.source_key)
        source = Path(task.source_path)
        saved_source = (
            settings.resolved_runtime() / "tasks" / ready.task_id / "source.bin"
        ).read_bytes()
        book = openpyxl.load_workbook(source)
        book.active["B2"] = 999
        book.save(source)
        model.calls.clear()
        rejected = service.run(task, retry=True, resume_candidate=True)
        assert rejected.error["code"] == "RESUME_SOURCE_CHANGED"
        assert (
            settings.resolved_runtime() / "tasks" / ready.task_id / "source.bin"
        ).read_bytes() == saved_source
        assert model.calls == []
        assert service.repo.current(task.source_key) is None
        with pytest.raises(AgentError, match="requires retry"):
            service.run(task, resume_candidate=True)
    finally:
        service.close()


def test_cli_resume_requires_saved_retry(capsys):
    from dc_agent.entrypoints.cli import main

    assert main(["run", "--example", "example", "--resume-candidate"]) == 2
    assert "INVALID_INPUT" in capsys.readouterr().out


def test_cli_without_command_shows_help_without_loading_settings_or_network(monkeypatch, capsys):
    from dc_agent.entrypoints import cli

    monkeypatch.setattr(cli, "Settings", lambda: pytest.fail("No command must not load settings"))
    assert cli.main([]) == 0
    output = capsys.readouterr().out
    assert "check-env" in output and "check-model" in output
    assert "arguments are required" not in output


@pytest.mark.parametrize("alter_name", [False, True])
def test_local_regex_repair_runs_real_source_and_keeps_semantic_gate(
    inputs, monkeypatch, alter_name
):
    from dc_agent.orchestration import workflow as service_module

    settings, task = inputs
    settings = settings.model_copy(update={"max_repairs": 0})
    task = task.model_copy(update={"title": "统计2026年3月"})
    pattern = r"(?<!\d)(20\d{2})年\s*(1[0-2]|[1-9])月(?!\d)".replace("\\", "\\\\")
    helper = "\nimport calendar\nimport re\ndef _date(title):\n"
    helper += f"    parts = re.findall({pattern!r}, title)\n"
    helper += '    if len(parts) != 1: return ""\n'
    helper += '    y, m = map(int, parts[0])\n    return f"{y:04d}-{m:02d}-{calendar.monthrange(y,m)[1]:02d}"\n'

    class DateModel(ScriptedModel):
        def complete(self, *args):
            reply = super().complete(*args)
            if args[0] == "design":
                reply.update(date_position="runtime:title", date_logic="dynamic month end")
            if args[0].startswith("generate"):
                reply["content"] = (
                    CODE.replace('"end_date": ""', '"end_date": _date(params.get("title", ""))')
                    + helper
                )
                reply["evidence"][0]["date_refs"] = ["runtime:title"]
            return reply

    if alter_name:
        proposer = service_module.propose_regex_escape_repair

        def unsafe_proposal(*args):
            proposal = proposer(*args)
            if proposal:
                from dc_agent.validation.candidate_repairs import RepairProposal

                candidate = proposal.candidate.model_copy(
                    update={
                        "content": proposal.candidate.content.replace(
                            'sheet["A2"].value', '"ChangedBusinessName"'
                        )
                    }
                )
                return RepairProposal(candidate, proposal.edits)

        monkeypatch.setattr(service_module, "propose_regex_escape_repair", unsafe_proposal)
    model = DateModel()
    service = make_service(settings, model)
    try:
        result = service.run(task)
        if alter_name:
            assert result.status == "review_required", result
            assert model.calls == ["design", "generate_0"]
            assert service.repo.current(task.source_key) is None
        else:
            assert result.status == "ready", result
            assert model.calls == ["design", "generate_0", "review_1"]
            workspace = settings.resolved_runtime() / "tasks" / result.task_id
            assert (workspace / "local-repair-1.json").exists()
            assert (workspace / "candidate-0.json").exists()
            report = json.loads((workspace / "validation-1.json").read_text(encoding="utf-8"))
            assert report["records"][0]["end_date"] == "2026-03-31"
    finally:
        service.close()


def test_semantic_feedback_survives_intermediate_execution_failure(inputs):
    settings, task = inputs
    settings = settings.model_copy(update={"max_repairs": 2})

    class RepairModel(ScriptedModel):
        def complete(self, *args):
            reply = super().complete(*args)
            operation, _, context, *_ = args
            if operation == "review_0":
                reply.update(approved=False, findings=["Check the actual parent label"])
            if operation == "generate_1":
                reply["content"] = CODE.replace('"ExcelData"', '"OUTDATA"')
            if operation == "generate_2":
                assert context["previous_semantic_findings"] == ["Check the actual parent label"]
                assert any(c["status"] == "fail" for c in context["validation_feedback"]["checks"])
            return reply

    model = RepairModel()
    service = make_service(settings, model)
    try:
        result = service.run(task)
        assert result.status == "ready", result
        assert model.calls == [
            "design",
            "generate_0",
            "review_0",
            "generate_1",
            "generate_2",
            "review_2",
        ]
    finally:
        service.close()


@pytest.mark.parametrize("explain_exclusion", [False, True])
def test_source_coverage_shrink_is_advisory_during_repair(inputs, explain_exclusion):
    settings, task = inputs
    source = Path(task.source_path)
    book = openpyxl.load_workbook(source)
    book.active.append(["表格样本量", 12])
    book.save(source)
    two_records = (
        CODE.replace('    return {"code":', '    result = {"code":')
        + """
    result["ExcelData"].append({"indicator_name": sheet["A3"].value,
        "indicator_data": str(sheet["B3"].value), "indicator_unit": "-", "end_date": ""})
    return result
"""
    )

    class CoverageModel(ScriptedModel):
        def complete(self, *args):
            reply = super().complete(*args)
            if args[0] == "generate_0":
                reply["content"] = two_records
                reply["evidence"].append(
                    {"record_index": 1, "name_refs": ["Sheet!A3"], "data_ref": "Sheet!B3"}
                )
            if args[0] == "review_0":
                reply.update(
                    approved=False, findings=["Check whether sample count is business data"]
                )
            if args[0] == "review_1":
                assert args[2]["coverage_comparison"]["missing_previous_data_refs"] == ["Sheet!B3"]
                reply["reviewed_refs"] += ["Sheet!A3", "Sheet!B3"]
                if explain_exclusion:
                    reply["excluded_source_refs"] = {
                        "Sheet!B3": "A3 labels table sample count, outside the requested business indicators"
                    }
            return reply

    service = make_service(settings, CoverageModel())
    try:
        result = service.run(task)
        assert result.status == "ready", result
        report = json.loads(
            (settings.resolved_runtime() / result.artifact["validation_path"]).read_text(
                encoding="utf-8"
            )
        )
        check = next(c for c in report["checks"] if c["name"] == "source_coverage_regression")
        assert check["status"] == ("pass" if explain_exclusion else "warning")
        assert not any(c["status"] == "fail" for c in report["checks"])
    finally:
        service.close()


def test_final_failure_exposes_parser_error_and_does_not_publish(inputs):
    settings, task = inputs
    settings = settings.model_copy(update={"max_repairs": 0})

    class MissingParameterModel(ScriptedModel):
        def complete(self, *args):
            reply = super().complete(*args)
            if args[0].startswith("generate"):
                reply["content"] = (
                    "def main(params):\n"
                    " return {'code': 1, 'msg': 'Missing fileServiceUrl', 'ExcelData': []}"
                )
            return reply

    model = MissingParameterModel()
    service = make_service(settings, model)
    try:
        result = service.run(task)
        assert result.status == "review_required"
        assert result.error["code"] == "VALIDATION_FAILED"
        assert "Missing fileServiceUrl" in result.error["message"]
        assert model.calls == ["design", "generate_0"]
        assert service.repo.current(task.source_key) is None
    finally:
        service.close()


def test_fenced_publication_and_source_lock(tmp_path):
    repo = Repository(tmp_path)
    try:
        token = repo.claim("source", 100)
        assert token
        assert repo.claim("source", 100) is None
        with pytest.raises(AgentError, match="expired"):
            repo.activate("source", "wrong-owner", {})
        repo.activate("source", token, {"version": "v1"})
        assert repo.current("source")["version"] == "v1"
        repo.release("source", "wrong-owner")
        assert repo.claim("source", 100) is None
        repo.release("source", token)
        assert repo.claim("source", 100)
    finally:
        repo.close()


def test_cache_corruption_is_not_silently_used(inputs):
    settings, task = inputs
    service = make_service(settings, ScriptedModel())
    try:
        result = service.run(task)
        service.assets.content_path(result.artifact).write_text("changed")
        failed = service.run(task.model_copy(update={"request_id": "next"}))
        assert failed.error["code"] == "ARTIFACT_CORRUPT"
        assert "checksum" in failed.error["message"]
    finally:
        service.close()


def test_live_gateway_mock_protocol_and_truncation():
    settings = Settings(
        _env_file=None,
        model_base_url="https://models.example/v1",
        model_name="test",
        model_api_key="not-a-real-key",
        model_retries=0,
    )

    def handler(request):
        assert request.url.path == "/v1/chat/completions"
        assert request.headers["authorization"] == "Bearer not-a-real-key"
        assert json.loads(request.content)["response_format"] == {"type": "json_object"}
        return httpx.Response(
            200,
            json={"choices": [{"finish_reason": "stop", "message": {"content": '{"ok":true}'}}]},
        )

    gateway = ChatCompletionsGateway(settings, httpx.MockTransport(handler))
    assert gateway.complete("test", "json", {}, {}, 10) == {"ok": True}
    bad = ChatCompletionsGateway(
        settings,
        httpx.MockTransport(
            lambda _: httpx.Response(
                200, json={"choices": [{"finish_reason": "length", "message": {"content": "{}"}}]}
            )
        ),
    )
    with pytest.raises(AgentError, match="finish"):
        bad.complete("test", "json", {}, {}, 10)


def test_missing_model_is_explicit(inputs):
    settings, task = inputs
    settings = settings.model_copy(update={"model_base_url": "", "model_name": ""})
    service = make_service(settings, ChatCompletionsGateway(settings))
    try:
        result = service.run(task)
        assert result.status == "failed"
        assert result.error["code"] == "MODEL_NOT_CONFIGURED"
        assert service.repo.current(task.source_key) is None
    finally:
        service.close()


def test_example_metadata_is_not_historical_output():
    task = load_task(example=Path(__file__).resolve().parents[1] / "example")
    assert task.fileType == "xlsx"
    assert len(task.GZKJG) == 21
    assert "excel_data" not in task.model_dump()
    assert "out_data" not in task.model_dump()
    assert task.GZKJG[0].rule_indicator_code.isdigit()


def test_task_rules_and_types(inputs):
    _, task = inputs
    assert task.source_key != task.model_copy(update={"fileType": "xls"}).source_key
    assert normalize_name("Beginning Stocks") == "Beginning Stocks"
    assert normalize_name("甲、乙,'") == "甲乙"
    for bad in ([], [{"rule_indicator_name": "empty"}], [{"rule_indicator_name": " "}]):
        with pytest.raises(ValueError):
            Task.model_validate({**task.model_dump(), "GZKJG": bad})
    mixed = Task.model_validate(
        {
            **task.model_dump(),
            "GZKJG": [{"rule_indicator_name": " "}, {"rule_indicator_name": "甲"}],
        }
    )
    assert len(mixed.GZKJG) == 1
    assert mixed.invalid_rule_rows == 1


def test_transient_failure_retry_reuses_completed_design(inputs):
    settings, task = inputs

    class InterruptedModel(ScriptedModel):
        failed_once = False

        def complete(self, operation, *args):
            if operation.startswith("generate") and not self.failed_once:
                self.failed_once = True
                raise AgentError("MODEL_TIMEOUT", "Simulated interrupted network", retryable=True)
            return super().complete(operation, *args)

    model = InterruptedModel()
    service = make_service(settings, model)
    try:
        assert service.run(task).error["code"] == "MODEL_TIMEOUT"
        assert service.run(task, retry=True).status == "ready"
        assert model.calls.count("design") == 1
    finally:
        service.close()


def test_prompt_branch_does_not_execute_target_model(tmp_path, write_catalog, write_reference):
    from dc_agent.validation import HEADINGS

    source = tmp_path / "source.md"
    source.write_text("单位：万元\n|指标|值|\n|甲|1|", encoding="utf-8")
    skill_root = tmp_path / "skills"
    skill = skill_root / "markdown-table-prompt-extraction" / "SKILL.md"
    skill.parent.mkdir(parents=True)
    skill.write_text("Markdown skill", encoding="utf-8")
    task = Task(
        pageTitle="site",
        fileType="md",
        markdownMD5="md-id",
        source_path=str(source),
        GZKJG=[{"rule_indicator_name": "甲", "rule_indicator_code": "123"}],
    )
    settings = Settings(
        _env_file=None,
        runtime_dir=tmp_path / "runtime",
        skill_root=skill_root,
        rule_dir=tmp_path / "rule",
    )
    write_catalog(settings.resolved_rules(), [[None, "site", "md-code", None, task.pageTitle]])
    write_reference(settings.resolved_rules() / "IndicatorName_Excel_Md.xlsx", code="other-code")

    class PromptModel:
        def complete(self, operation, system, context, schema, timeout):
            if operation == "preflight":
                return {"verdict": "has_data", "reason": "Markdown has a real indicator table"}
            if operation == "design":
                return {
                    "summary": "Markdown",
                    "name_logic": "指标行",
                    "unit_logic": "单位行",
                    "unit_position": "L1",
                    "date_logic": "空",
                    "date_position": "未提供",
                    "covered_sections": ["markdown"],
                }
            if operation.startswith("generate"):
                content = "<!-- 来源：site；source_rules_v2 -->\n"
                content += "\n".join(
                    f"# {heading}\n具体规则：源Markdown候选精确唯一匹配GZK.GZKJG.rule_indicator_name。"
                    for heading in HEADINGS
                )
                content += '\n{"out_data":[{"indicator_code":"rule_indicator_code","indicator_name":"","indicator_data":"","indicator_unit":"","end_date":""}]}'
                return {"content": content}
            return {"approved": True, "findings": [], "reviewed_refs": ["L1", "L3"]}

    class NoExecution:
        def run(self, *args):
            pytest.fail("Prompt branch must not run a target extraction model or Python")

    service = AssetService(
        settings,
        PromptModel(),
        FileSourceProvider(settings),
        LocalReferences(),
        NoExecution(),
        SkillRegistry(skill_root),
    )
    try:
        result = service.run(task)
        assert result.status == "ready", result
        report = json.loads(
            (settings.resolved_runtime() / result.artifact["validation_path"]).read_text(
                encoding="utf-8"
            )
        )
        assert (
            next(c for c in report["checks"] if c["name"] == "target_model_execution")["status"]
            == "not_run"
        )
    finally:
        service.close()


def test_task_deadline_failure_does_not_publish(inputs):
    settings, task = inputs
    settings = settings.model_copy(update={"task_timeout": 0.000001})
    service = make_service(settings, ScriptedModel())
    try:
        assert service.run(task).error["code"] == "BUDGET_EXCEEDED"
        assert service.repo.current(task.source_key) is None
    finally:
        service.close()


def test_cli_retry_by_saved_task_id_preserves_request_and_source(inputs, monkeypatch):
    from dc_agent.domain import Result
    from dc_agent.entrypoints import cli

    settings, task = inputs
    repo = Repository(settings.resolved_runtime())
    row = repo.register(task)
    repo.close()
    observed = []

    class Service:
        def run(self, saved_task, *, retry, rebuild):
            observed.append((saved_task, retry, rebuild))
            return Result(status="failed", task_id=row["id"], source_key=saved_task.source_key)

        def close(self):
            pass

    monkeypatch.setattr(cli, "Settings", lambda: settings)
    monkeypatch.setattr(cli, "build_service", lambda _: Service())
    assert cli.main(["run", "--task-id", row["id"], "--retry"]) == 3
    saved_task, retry, rebuild = observed[0]
    assert saved_task.request_id == task.request_id
    assert saved_task.source_path == task.source_path
    assert saved_task.input_hash == task.input_hash
    assert retry and not rebuild


def test_no_data_stops_before_skill_design_generation_or_execution(inputs):
    settings, task = inputs
    settings = settings.model_copy(update={"skill_root": Path(task.source_path).parent / "missing"})

    class EmptyModel:
        calls = []

        def complete(self, operation, system, context, schema, timeout):
            self.calls.append(operation)
            assert "GZKJG" not in context
            return {"verdict": "no_data", "reason": "Explicitly no source data"}

    model = EmptyModel()
    service = make_service(settings, model)
    try:
        result = service.run(task)
        assert result.status == "failed"
        assert result.error["code"] == "FilenoData"
        assert result.error["stage"] == "preflight"
        assert result.error["retryable"] is False
        assert result.preflight["allowed"] is False
        assert model.calls == ["preflight"]
        assert service.repo.current(task.source_key) is None
        assert not list(settings.resolved_code().rglob("*.py"))
    finally:
        service.close()


@pytest.mark.parametrize("verdict", ["uncertain", "has_data"])
def test_preflight_is_before_design_and_uncertain_still_generates(inputs, verdict):
    settings, task = inputs
    observed = []

    class OrderedModel(ScriptedModel):
        def complete(self, operation, *args):
            observed.append(operation)
            if operation == "preflight":
                return {"verdict": verdict, "reason": "Coarse judgment"}
            return super().complete(operation, *args)

    service = make_service(settings, OrderedModel())
    try:
        result = service.run(task)
        assert result.status == "ready", result
        assert observed == ["preflight", "design", "generate_0", "review_0"]
        assert result.preflight["allowed"]
        assert result.artifact["content_path"] == "excel/xlsx/站点_统计.py"
    finally:
        service.close()


def test_preflight_failure_does_not_turn_into_no_data_or_generate(inputs):
    settings, task = inputs
    calls = []

    class FailingModel:
        def complete(self, operation, *args):
            calls.append(operation)
            raise AgentError("MODEL_TIMEOUT", "Simulated preflight failure", retryable=True)

    service = make_service(settings, FailingModel())
    try:
        result = service.run(task)
        assert result.error["code"] == "MODEL_TIMEOUT"
        assert result.error["stage"] == "preflight"
        assert calls == ["preflight"]
        assert service.repo.current(task.source_key) is None
    finally:
        service.close()


def test_cache_hit_skips_source_and_all_model_operations(inputs):
    settings, task = inputs
    model = ScriptedModel()
    service = make_service(settings, model)
    try:
        ready = service.run(task)
        assert ready.status == "ready"

        class NoRead:
            def read(self, task):
                pytest.fail("Cache reuse must not download source")

        service.source = NoRead()
        cached = service.run(task.model_copy(update={"request_id": "cache-without-source"}))
        assert cached.status == "ready" and cached.cache_hit
        assert cached.preflight["status"] == "not_run"
        assert model.calls == ["design", "generate_0", "review_0"]
        assert model.preflight_calls == 1
    finally:
        service.close()


def test_failed_no_data_rebuild_keeps_the_activated_asset(inputs):
    settings, task = inputs

    class Model(ScriptedModel):
        no_data = False

        def complete(self, operation, *args):
            if operation == "preflight" and self.no_data:
                return {"verdict": "no_data", "reason": "Explicit no data"}
            return super().complete(operation, *args)

    model = Model()
    service = make_service(settings, model)
    try:
        first = service.run(task)
        assert first.status == "ready"
        model.no_data = True
        rebuilt = service.run(task.model_copy(update={"request_id": "new-no-data"}), rebuild=True)
        assert rebuilt.error["code"] == "FilenoData"
        assert service.repo.current(task.source_key) == first.artifact
        assert service.assets.lookup(task) == first.artifact
    finally:
        service.close()


def test_idempotent_old_result_points_to_its_immutable_version(inputs):
    settings, task = inputs
    service = make_service(settings, ScriptedModel())
    try:
        first = service.run(task)
        assert first.status == "ready"
        second = service.run(task.model_copy(update={"request_id": "rebuild"}), rebuild=True)
        assert second.status == "ready"
        assert second.artifact["version"] != first.artifact["version"]
        replay = service.run(task)
        assert replay.artifact["version"] == first.artifact["version"]
        assert replay.artifact["content_path"] == first.artifact["version_path"]
        assert service.assets.content_path(replay.artifact).is_file()
    finally:
        service.close()
