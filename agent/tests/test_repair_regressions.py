"""Repair edge cases use local workbook bytes and scripted model replies only."""

import json

import openpyxl
import pytest
from fastapi.testclient import TestClient
from test_application import CODE, ScriptedModel, make_service
from test_application import inputs as inputs
from test_http_integration import AUTH, PAYLOAD, TOKEN
from test_workflow_nodes import inputs as node_inputs
from test_workflow_nodes import load_node

from dc_agent.config import Settings
from dc_agent.domain import Candidate, Task
from dc_agent.entrypoints.http import create_app
from dc_agent.repair.session import RepairSession
from dc_agent.validation import Validator


def read(path):
    return json.loads(path.read_text(encoding="utf-8"))


@pytest.mark.parametrize("kind", ["absent", "null", "json-null"])
def test_missing_rules_have_same_terminal_code_in_http_and_platform(tmp_path, monkeypatch, kind):
    payload = dict(PAYLOAD)
    business = node_inputs()
    if kind == "absent":
        payload.pop("GZKJG")
        business.pop("GZKJG")
    else:
        payload["GZKJG"] = business["GZKJG"] = None if kind == "null" else "null"
    settings = Settings(_env_file=None, runtime_dir=tmp_path / "runtime", api_token=TOKEN)
    app = create_app(settings, start_worker=False)
    with TestClient(app) as client:
        response = client.post("/v3/excel/jobs", json=payload, headers=AUTH)
    assert response.status_code == 422
    assert response.json()["error_code"] == "without_GZKJG"
    with app.state.store._connection() as db:
        assert db.execute("SELECT COUNT(*) FROM jobs").fetchone()[0] == 0

    node = load_node("外部框架接入.py", "repair_regression_empty_rules")
    monkeypatch.setattr(node.requests, "Session", lambda: pytest.fail("Unexpected HTTP call"))
    record = node.main(business)["agent"][0]
    assert record["status"] == "invalid_input"
    assert record["error_code"] == "without_GZKJG"
    assert set(record) == set(node.AGENT_FIELDS)
    assert all(isinstance(value, str) for value in record.values())
    assert all(record[field] == "" for field in node._CODE_FIELDS)


def test_mixed_invalid_rules_preserve_valid_rule_without_rejecting_job(tmp_path):
    rules = [
        {"rule_indicator_name": "empty", "rule_indicator_code": 999},
        {"rule_indicator_name": None},
        {"rule_indicator_name": False},
        {"rule_indicator_name": "指标甲", "rule_indicator_code": 123},
    ]
    node = load_node("外部框架接入.py", "repair_regression_mixed_rules")
    assert node._payload(node_inputs(GZKJG=rules))["GZKJG"] == rules
    settings = Settings(_env_file=None, runtime_dir=tmp_path / "runtime", api_token=TOKEN)
    app = create_app(settings, start_worker=False)
    with TestClient(app) as client:
        response = client.post("/v3/excel/jobs", json={**PAYLOAD, "GZKJG": rules}, headers=AUTH)
    assert response.status_code == 202
    saved = app.state.store.get(response.json()["job_id"])["payload"]
    assert saved["invalid_rule_rows"] == 3
    assert saved["GZKJG"] == [{"rule_indicator_name": "指标甲", "rule_indicator_code": "123"}]


class FixedResult:
    def __init__(self, end_date):
        self.end_date = end_date

    def run(self, *args, **kwargs):
        return {
            "code": 0,
            "msg": "success",
            "ExcelData": [
                {
                    "indicator_name": "指标甲",
                    "indicator_data": "0",
                    "indicator_unit": "-",
                    "end_date": self.end_date,
                }
            ],
        }


def validation_fixture(*, index=0, date_refs=()):
    task = Task.model_validate({**PAYLOAD, "title": "2026年3月统计"})
    candidate = Candidate(
        content='def main(params): return {"ExcelData": []}',
        evidence=[
            {
                "record_index": index,
                "name_refs": ["Sheet!A2"],
                "data_ref": "Sheet!B2",
                "date_refs": list(date_refs),
            }
        ],
    )
    inspection = {
        "sha256": "synthetic-source",
        "format": "xlsx",
        "sections": [
            {
                "id": "Sheet",
                "cells": [{"ref": "Sheet!A2", "value": "指标甲"}, {"ref": "Sheet!B2", "value": 0}],
            }
        ],
    }
    return task, candidate, inspection


@pytest.mark.parametrize("value", [None, False, 0])
def test_non_string_empty_dates_are_contract_failures(value):
    task, candidate, inspection = validation_fixture()
    report = Validator(FixedResult(value)).validate(task, candidate, inspection, b"local")
    assert not report.passed
    check = next(check for check in report.checks if check.name == "runtime_contract")
    assert check.status == "fail"
    assert "end_date" in check.message


def test_evidence_outside_output_becomes_repair_diagnostic_not_index_error(tmp_path):
    task, candidate, inspection = validation_fixture(index=10, date_refs=["runtime:title"])
    report = Validator(FixedResult("")).validate(task, candidate, inspection, b"local")
    assert not report.passed
    evidence = next(check for check in report.checks if check.name == "source_evidence")
    assert evidence.status == "fail" and "Evidence index outside output" in evidence.message
    session = RepairSession(tmp_path, tmp_path, task, inspection, "regression-policy")
    session.validation(report, candidate.model_dump(), 0)
    plan = session.plan(report, candidate.model_dump())
    assert plan["targets"]
    assert plan["record_fragments"] == []


def test_literal_empty_date_is_repaired_against_runtime_title(inputs):
    settings, task = inputs
    task = task.model_copy(update={"title": "2026年3月统计"})
    fixed = (
        CODE.replace('"end_date": ""', '"end_date": _month_end(params["title"])')
        + r"""
import calendar
import re
def _month_end(title):
    match = re.search(r"(20\d{2})年(\d{1,2})月", title)
    if not match:
        return ""
    year, month = map(int, match.groups())
    return f"{year:04d}-{month:02d}-{calendar.monthrange(year, month)[1]:02d}"
"""
    )

    class EmptyDate(ScriptedModel):
        def complete(self, operation, *args):
            reply = super().complete(operation, *args)
            if operation == "generate_0":
                reply["content"] = CODE.replace('"end_date": ""', '"end_date": "empty"')
            elif operation == "generate_1":
                reply["content"] = fixed
                reply["evidence"][0]["date_refs"] = ["runtime:title"]
            return reply

    service = make_service(settings, EmptyDate())
    try:
        result = service.run(task)
        assert result.status == "ready", result
        workspace = settings.resolved_runtime() / "tasks" / result.task_id
        baseline = read(workspace / "validation-0.json")
        failed = next(c for c in baseline["checks"] if c["name"] == "runtime_contract")
        assert failed["status"] == "fail"
        assert failed["details"]["category"] == "end_date_empty"
        selected = read(settings.resolved_runtime() / result.artifact["validation_path"])
        assert selected["records"][0]["end_date"] == "2026-03-31"
    finally:
        service.close()


def test_unit_advisory_mixed_with_real_name_error_still_requires_repair(inputs):
    settings, task = inputs

    class MixedReview(ScriptedModel):
        def complete(self, operation, *args):
            reply = super().complete(operation, *args)
            if operation == "generate_0":
                reply["content"] = CODE.replace('sheet["A2"].value', '"错误业务名称"')
            if operation == "review_0":
                reply.update(
                    approved=False,
                    findings=[
                        "Source unit differs from KB",
                        "Name contradicts actual source label",
                    ],
                    issues=[
                        {
                            "category": "unit_conversion",
                            "message": "Source unit differs from KB",
                            "record_indices": [0],
                            "source_refs": ["Sheet!B2"],
                        },
                        {
                            "category": "name_semantics",
                            "message": "Name contradicts actual source label",
                            "record_indices": [0],
                            "source_refs": ["Sheet!A2"],
                        },
                    ],
                )
            return reply

    model = MixedReview()
    service = make_service(settings, model)
    try:
        result = service.run(task)
        assert result.status == "ready", result
        assert "generate_1" in model.calls
        workspace = settings.resolved_runtime() / "tasks" / result.task_id
        baseline = read(workspace / "validation-0.json")
        review = next(c for c in baseline["checks"] if c["name"] == "model_semantic_review")
        assert review["status"] == "fail"
        selected = read(settings.resolved_runtime() / result.artifact["validation_path"])
        assert selected["records"][0]["indicator_name"] == "指标甲"
    finally:
        service.close()


def test_source_period_columns_are_not_overridden_by_different_runtime_title(inputs):
    settings, task = inputs
    task = task.model_copy(update={"title": "2026年3月发布的历史统计"})
    book = openpyxl.load_workbook(task.source_path)
    sheet = book.active
    sheet["B1"], sheet["C1"] = "2026年1月", "2026年2月"
    sheet["B2"], sheet["C2"] = 10, 20
    book.save(task.source_path)
    book.close()
    book = openpyxl.load_workbook(task.source_path, data_only=True)
    assert list(book.active.values) == [("指标", "2026年1月", "2026年2月"), ("指标甲", 10, 20)]
    book.close()
    source_code = r"""
import calendar
import re
from io import BytesIO
import requests
import openpyxl
def main(params):
    raw = requests.get("http://service-1.example.invalid/api/file/download/" + params["fileMd5"]).content
    sheet = openpyxl.load_workbook(BytesIO(raw), data_only=True).active
    records = []
    for column in (2, 3):
        match = re.search(r"(20\d{2})年(\d{1,2})月", sheet.cell(1, column).value)
        year, month = map(int, match.groups())
        records.append({"indicator_name": sheet["A2"].value,
                        "indicator_data": sheet.cell(2, column).value,
                        "indicator_unit": "-",
                        "end_date": f"{year:04d}-{month:02d}-{calendar.monthrange(year, month)[1]:02d}"})
    return {"code": 0, "msg": "success", "ExcelData": records}
"""

    class PeriodColumns(ScriptedModel):
        def complete(self, operation, *args):
            reply = super().complete(operation, *args)
            if operation == "design":
                reply.update(
                    date_logic="Use each actual data column month",
                    date_position="Sheet!B1 and Sheet!C1",
                )
            elif operation == "generate_0":
                reply["content"] = source_code
                reply["evidence"] = [
                    {
                        "record_index": index,
                        "name_refs": ["Sheet!A2"],
                        "data_ref": f"Sheet!{column}2",
                        "date_refs": [f"Sheet!{column}1"],
                    }
                    for index, column in enumerate(("B", "C"))
                ]
            return reply

    model = PeriodColumns()
    service = make_service(settings, model)
    try:
        result = service.run(task)
        assert result.status == "ready", result
        assert model.calls == ["design", "generate_0", "review_0"]
        selected = read(settings.resolved_runtime() / result.artifact["validation_path"])
        assert [row["end_date"] for row in selected["records"]] == ["2026-01-31", "2026-02-28"]
        assert [row["indicator_data"] for row in selected["records"]] == [10, 20]
    finally:
        service.close()
