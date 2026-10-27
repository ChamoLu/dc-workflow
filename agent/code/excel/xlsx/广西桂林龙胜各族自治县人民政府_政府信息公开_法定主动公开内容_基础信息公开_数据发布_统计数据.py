# 数据源：广西桂林龙胜各族自治县人民政府
# 数据源路径代码：2087036847535702018；数据来源编码：01020002334
# 目标表：usrDFJJYXSJB
# 名称参考：IndicatorName_Excel_Md.xlsx / 2087036847535702018，A1:B20、K1:K20。
# 参考仅用于审查，不作为运行时输入或白名单。
# 源事实修正：D3中的%是单位，而不是整个“同比增长%”作为单位。
# A6:A8为A5的产业子项，名称保留地区生产总值父级，不照搬参考的扁平名称。
# A15明确部分数据为季度指标；A1只给月度表期，未明确季度期间。
# 对这些季度指标日期留空，不擅自赋予五月末或推算一季度日期。
# 其他指标日期取实际下载文件A1，优先于外部title；不使用publictime。

import calendar
import io
import math
import re
import unicodedata

import requests
import openpyxl


def _text(value):
    if value is None:
        return ''
    value = unicodedata.normalize('NFKC', str(value))
    return value.replace('\n', '').replace('\r', '').replace('、', '').replace(',', '').replace("'", '').strip()


def _label(value):
    text = _text(value)
    return re.sub(r'^[一二三四五六七八九十百]+[、.．]?\s*', '', text)


def _period(value):
    text = _text(value)
    match = re.search(r'(\d{4})年\s*(\d{1,2})(?:\s*[-—至~]\s*(\d{1,2}))?月', text)
    if not match:
        return ''
    year = int(match.group(1))
    month = int(match.group(3) or match.group(2))
    if not 1 <= month <= 12 or not 1 <= year <= 9999:
        return ''
    return '%04d-%02d-%02d' % (year, month, calendar.monthrange(year, month)[1])


def _numeric(value):
    if isinstance(value, bool):
        return None, ''
    if isinstance(value, (int, float)):
        if math.isfinite(value):
            return value, ''
        return None, ''
    text = _text(value)
    annotation = re.fullmatch(r'(?:提高|上升|增加)\s*(\d+(?:\.\d+)?)\s*(个百分点|百分点)', text)
    if annotation:
        return annotation.group(1), annotation.group(2)
    if re.fullmatch(r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)', text):
        return text, ''
    return None, ''


def _extract(params):
    md5 = params.get('fileMd5')
    if not md5:
        raise ValueError('缺少fileMd5')
    response = requests.get(
        'http://service-1.example.invalid/api/file/download/' + str(md5),
        timeout=60,
    )
    response.raise_for_status()
    workbook = openpyxl.load_workbook(io.BytesIO(response.content), data_only=True)
    records = []
    evidence = []
    try:
        if '主要指标' not in workbook.sheetnames:
            raise ValueError('未找到主要指标Sheet')
        sheet = workbook['主要指标']
        if _text(sheet['C3'].value) != '总量' or _text(sheet['D3'].value) != '同比增长%':
            raise ValueError('源表列结构发生变化，需复核')
        monthly_date = _period(sheet['A1'].value)
        parent = ''
        parent_ref = ''
        parent_quarterly = False
        for row in range(5, sheet.max_row + 1):
            raw = sheet.cell(row, 1).value
            label = _text(raw)
            if not label:
                continue
            if label.startswith(('注:', '注：')):
                break
            quarterly = '(季度数)' in label
            name = _label(raw).replace('(季度数)', '').strip()
            if not name:
                continue
            refs = ['主要指标!A%d' % row]
            is_child = isinstance(raw, str) and bool(re.match(r'^\s+', raw))
            if is_child and parent:
                name = parent + name
                refs.insert(0, parent_ref)
                quarterly = quarterly or parent_quarterly
            else:
                parent = name
                parent_ref = '主要指标!A%d' % row
                parent_quarterly = quarterly
            # 季度频率由行标签及A15说明支持，但具体季度未给出。
            end_date = '' if quarterly else monthly_date
            date_refs = [] if quarterly or not monthly_date else ['主要指标!A1']
            for column in (3, 4):
                cell = sheet.cell(row, column)
                value, annotation_unit = _numeric(cell.value)
                if value is None:
                    continue
                header = sheet.cell(3, column)
                if column == 4:
                    unit = '%'
                    unit_refs = ['主要指标!D3']
                else:
                    unit = _text(sheet.cell(row, 2).value) or '-'
                    unit_refs = ['主要指标!B%d' % row] if unit != '-' else []
                if annotation_unit:
                    unit = annotation_unit
                    unit_refs = ['主要指标!' + cell.coordinate]
                records.append({
                    'indicator_name': name + _text(header.value),
                    'indicator_data': value,
                    'indicator_unit': unit,
                    'end_date': end_date,
                })
                evidence.append({
                    'record_index': len(records) - 1,
                    'name_refs': refs + ['主要指标!' + header.coordinate],
                    'data_ref': '主要指标!' + cell.coordinate,
                    'unit_refs': unit_refs,
                    'date_refs': list(date_refs),
                })
    finally:
        workbook.close()
    return {'code': 0, 'msg': 'success', 'ExcelData': records}, evidence


def main(params):
    result, _ = _extract(params)
    return result


def _dc_validate(params):
    result, evidence = _extract(params)
    result['_dc_evidence'] = evidence
    return result
