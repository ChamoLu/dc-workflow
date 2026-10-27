# 数据源：上街区人民政府
# 数据源路径代码：2085711403154546690；数据来源编码：01020002879
# 路径：上街区人民政府\政务公开\政府信息公开\法定主动公开内容\统计信息\统计月报
# 目标表：usrDFJJYXSJB
# 规则来源：rule/todo.xlsx / Sheet1第5行；名称参考精确Sheet 2085711403154546690。
# 单位逻辑：读取本表单位声明、真实行名括号、增长表头及CPI百分点说明；无来源为-。
# 单位位置：各块单位声明/实际指标行括号/比去年同期±%表头，不按固定行号取单位。
# 截止日期位置：运行时外部title；源内无具体年月，唯一统计年月归月末，未知则空。
# 按表题、表头及真实行标签定位四块，支持新增收入行及表头下移，不将标题/脚注作为数据。
# 三次产业结构复合比例不属于单数值；不拆造，不换算，不去重，不四舍五入。
# GZK仅用于设计审查；生产main不读取GZK、参考工作簿或本地文件。

import calendar
import math
import re
import unicodedata
from decimal import Decimal, InvalidOperation
from io import BytesIO

import openpyxl
import requests

_DOWNLOAD_URL = "http://service-1.example.invalid/api/file/download/"
_TITLES = {
    "全区主要经济指标": "summary",
    "规模以上工业增加值": "industry",
    "固定资产投资完成额": "investment",
    "社会消费品零售总额": "retail",
}
_UNITS = {"万元", "元", "%", "吨标准煤", "吨标准煤/万元"}


def _text(value):
    if value is None:
        return ""
    value = unicodedata.normalize("NFKC", str(value))
    return re.sub(r"\s+", "", value).replace("、", "").replace(",", "").replace("'", "")


def _number(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value if not isinstance(value, float) or math.isfinite(value) else None
    if not isinstance(value, str):
        return None
    cleaned = _text(value)
    if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", cleaned):
        return None
    try:
        return cleaned if Decimal(cleaned).is_finite() else None
    except InvalidOperation:
        return None


def _unit(label, default):
    units = [unit for unit in re.findall(r"\(([^()]*)\)", label) if unit in _UNITS]
    return units[-1] if units else default


def _end_date(title):
    periods = re.findall(r"(?<!\d)((?:19|20)\d{2})年(1[0-2]|0?[1-9])月", _text(title))
    if len(periods) != 1:
        return ""
    year, month = map(int, periods[0])
    return "%04d-%02d-%02d" % (year, month, calendar.monthrange(year, month)[1])


def _ref(sheet, row, column):
    return "%s!%s%d" % (sheet.title, openpyxl.utils.get_column_letter(column), row)


def _blocks(sheet):
    result = []
    for row in sheet.iter_rows():
        cells = [cell for cell in row if cell.value is not None]
        labels = {cell.column: _text(cell.value) for cell in cells}
        name_columns = [column for column, label in labels.items() if label == "指标"]
        growth_columns = [column for column, label in labels.items() if label == "比去年同期±%"]
        if not name_columns or not growth_columns:
            continue
        cumulative_columns = [column for column, label in labels.items() if label == "累计"]
        if len(name_columns) != 1 or len(growth_columns) != 1 or len(cumulative_columns) > 1:
            raise ValueError("ambiguous indicator/data header")
        heading = cells[0].row
        title_cell = None
        for preceding in range(heading - 1, 0, -1):
            values = [cell for cell in sheet[preceding] if cell.value is not None]
            if len(values) == 1 and _text(values[0].value) in _TITLES:
                title_cell = values[0]
                break
        if title_cell is None:
            raise ValueError("missing recognized section title before header")
        default_unit, unit_ref = "-", []
        for preceding in range(title_cell.row + 1, heading):
            for cell in sheet[preceding]:
                match = re.fullmatch(r"单位[:：](.+)", _text(cell.value))
                if match:
                    default_unit = match.group(1)
                    unit_ref = [_ref(sheet, cell.row, cell.column)]
        result.append(
            {
                "row": heading,
                "label_column": name_columns[0],
                "growth_column": growth_columns[0],
                "cumulative_column": cumulative_columns[0] if cumulative_columns else None,
                "title": _text(title_cell.value),
                "title_ref": _ref(sheet, title_cell.row, title_cell.column),
                "kind": _TITLES[_text(title_cell.value)],
                "default_unit": default_unit,
                "unit_refs": unit_ref,
            }
        )
    if [item["kind"] for item in result] != ["summary", "industry", "investment", "retail"]:
        raise ValueError("expected four recognized statistical sections")
    return result


def _extract(sheet, title):
    records, evidence = [], []
    end_date = _end_date(title)
    date_refs = ["runtime:title"] if end_date else []
    blocks = _blocks(sheet)

    def append(name, data, unit, name_refs, data_ref, unit_refs):
        index = len(records)
        records.append(
            {
                "indicator_name": name,
                "indicator_data": data,
                "indicator_unit": unit,
                "end_date": end_date,
            }
        )
        evidence.append(
            {
                "record_index": index,
                "name_refs": name_refs,
                "data_ref": data_ref,
                "unit_refs": unit_refs,
                "date_refs": date_refs,
            }
        )

    for index, block in enumerate(blocks):
        end_row = blocks[index + 1]["row"] - 1 if index + 1 < len(blocks) else sheet.max_row
        parent_label, parent_ref = "", ""
        group_label, group_ref = "", ""
        investment_total_ref = ""
        note_refs = [
            _ref(sheet, row, cell.column)
            for row in range(block["row"] + 1, end_row + 1)
            for cell in sheet[row]
            if cell.value is not None and "带*为比年初增长" in _text(cell.value)
        ]
        for row in range(block["row"] + 1, end_row + 1):
            raw_label = sheet.cell(row, block["label_column"]).value
            label = _text(raw_label)
            if not label or label.startswith("注") or label in _TITLES or label == "三次产业结构":
                # Section totals with data are handled below; pure title rows are not data.
                if label not in _TITLES or not any(
                    _number(sheet.cell(row, column).value) is not None
                    for column in (block["cumulative_column"], block["growth_column"])
                    if column is not None
                ):
                    continue
            label_ref = _ref(sheet, row, block["label_column"])
            marked = label.startswith("#")
            label = label.lstrip("#").replace("*", "")
            name, name_refs = label, [label_ref]
            if block["kind"] == "industry":
                name_refs.insert(0, block["title_ref"])
                if label.startswith("占规上工业比重"):
                    if not parent_label or not parent_label.endswith("产业"):
                        raise ValueError("industry share row has no source parent")
                    name = block["title"] + parent_label + label
                    name_refs.insert(1, parent_ref)
                elif label == "工业增加值":
                    name = block["title"]
                else:
                    if label == "工业主导产业合计":
                        group_label, group_ref = label, label_ref
                    elif marked:
                        group_label, group_ref = "", ""
                    name = block["title"] + (group_label if group_label != label else "") + label
                    if group_label and group_label != label:
                        name_refs.insert(1, group_ref)
                    parent_label, parent_ref = label, label_ref
            elif block["kind"] == "investment":
                if label.startswith(("按产业类型分", "按建设性质分")):
                    group_label, group_ref = label, label_ref
                    continue
                if label == "固定资产投资":
                    investment_total_ref = label_ref
                name = block["title"]
                name_refs.insert(0, block["title_ref"])
                if label != "固定资产投资":
                    if not investment_total_ref:
                        raise ValueError("investment detail has no source total")
                    name += "固定资产投资" + group_label
                    name_refs.append(investment_total_ref)
                    if group_label:
                        name_refs.append(group_ref)
                name += label
            elif block["kind"] == "retail":
                if label.startswith(("按规模分", "按行业分")):
                    group_label, group_ref = label, label_ref
                    continue
                name = (
                    block["title"]
                    if label == block["title"]
                    else block["title"] + group_label + label
                )
                name_refs.insert(0, block["title_ref"])
                if group_label:
                    name_refs.insert(1, group_ref)
            for column, cumulative in (
                (block["cumulative_column"], True),
                (block["growth_column"], False),
            ):
                if column is None:
                    continue
                value = sheet.cell(row, column).value
                data = _number(value)
                data_ref = _ref(sheet, row, column)
                header_ref = _ref(sheet, block["row"], column)
                column_name = _text(sheet.cell(block["row"], column).value)
                if data is None:
                    match = (
                        re.fullmatch(
                            r"(?:提高|上升|增加)(\+?\d+(?:\.\d+)?)(个百分点|百分点)", _text(value)
                        )
                        if not cumulative and "居民消费价格总指数" in label
                        else None
                    )
                    if match:
                        append(
                            name + column_name,
                            match.group(1).lstrip("+"),
                            match.group(2),
                            name_refs + [header_ref],
                            data_ref,
                            [data_ref],
                        )
                    continue
                if cumulative:
                    unit = _unit(label, block["default_unit"])
                    unit_refs = [label_ref] if _unit(label, "") else block["unit_refs"]
                else:
                    unit, unit_refs = "%", [header_ref]
                    if "*" in _text(raw_label) and note_refs:
                        column_name = "比年初增长"
                append(
                    name + column_name,
                    data,
                    unit,
                    name_refs + [header_ref] + (note_refs if column_name == "比年初增长" else []),
                    data_ref,
                    unit_refs,
                )
    return records, evidence


def _run(params, traced):
    file_md5 = params.get("fileMd5")
    if not isinstance(file_md5, str) or not file_md5.strip():
        raise ValueError("fileMd5 is required")
    response = requests.get(_DOWNLOAD_URL + file_md5, timeout=60)
    response.raise_for_status()
    workbook = openpyxl.load_workbook(BytesIO(response.content), data_only=True)
    try:
        sheets = [sheet for sheet in workbook if _text(sheet.title) == "主要经济指标"]
        if len(sheets) != 1:
            raise ValueError("expected unique sheet 主要经济指标")
        records, evidence = _extract(sheets[0], params.get("title", ""))
    finally:
        workbook.close()
    result = {"code": 0, "msg": "success", "ExcelData": records}
    if traced:
        result["_dc_evidence"] = evidence
    return result


def main(params):
    return _run(params, False)


def _dc_validate(params):
    return _run(params, True)
