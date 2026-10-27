# 数据源：方山县人民政府；数据来源编码：01020002178
# 数据源路径代码：2082491763079131138；目标表：usrDFJJYXSJB
# 设计依据：实际Sheet2、Sheet3、有及各自表题，不使用外部title覆盖表内期间。
# 名称参考：IndicatorName_Excel_Md.xlsx / 2082491763079131138，A1:K18。
# 参考仅供设计核对，不作为运行时输入或指标白名单。
# 冲突处理：任务所述Sheet1不存在；同比列的显式%优先于行物量单位。
# 有!B10明确交通运输业为错月数据，但未明确月份，本版本暂不提取该行。
# 覆盖建议：去年同期、全市增速、排名及辅助计算列未提取；不默认去重。
# 当前环境未执行真实bytes回归，以下代码不声称已完成运行验证。
import calendar
import io
import math
import re
import unicodedata

import requests
import openpyxl


DOWNLOAD_BASE = 'http://service-1.example.invalid/api/file/download/'


def clean(value):
    if value is None:
        return ''
    text = unicodedata.normalize('NFKC', str(value))
    return text.replace('\n', '').replace('\r', '').replace('、', '').replace(',', '').replace("'", '').strip()


def numeric(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value if math.isfinite(value) else None
    text = clean(value)
    if re.fullmatch(r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)', text):
        return text
    return None


def period(value):
    text = clean(value)
    match = re.search(r'(\d{4})年\s*(?:\d+\s*[-—至]\s*)?(\d{1,2})月', text)
    if not match:
        return ''
    year, month = map(int, match.groups())
    if not 1 <= month <= 12:
        return ''
    return '%04d-%02d-%02d' % (year, month, calendar.monthrange(year, month)[1])


def label(value):
    # 序号不属于业务名称；仅清除标签开头的中文章节序号。
    return re.sub(r'^[一二三四五六七八九十]+[、.．]', '', unicodedata.normalize('NFKC', str(value or '')).strip()).replace('、', '').replace(',', '').replace("'", '').replace('\n', '').replace('\r', '').strip()


def extract(book):
    records = []
    evidence = []

    def block(sheet_name, title_row, header_row, rows, parents=None, category_row=None):
        if sheet_name not in book.sheetnames:
            return
        sheet = book[sheet_name]
        parents = parents or {}
        date = period(sheet.cell(title_row, 2).value)

        def ref(row, col):
            return sheet.cell(row, col).coordinate.join([sheet_name + '!', ''])

        for row in rows:
            row_label = label(sheet.cell(row, 2).value)
            if not row_label:
                continue
            parent_rows = list(parents.get(row, []))
            if category_row is not None:
                parent_rows = [category_row] + parent_rows
            name_parts = [label(sheet.cell(p, 2).value) for p in parent_rows]
            name_parts.append(row_label)
            for col in (4, 6):
                value = sheet.cell(row, col).value
                number = numeric(value)
                annotation = re.fullmatch(r'(?:提高|上升|增加)\s*(\d+(?:\.\d+)?)\s*(个百分点|百分点)', clean(value))
                if number is None and annotation is None:
                    continue
                header = clean(sheet.cell(header_row, col).value)
                if not header:
                    continue
                suffix = re.sub(r'1-\d+月份?', '', header)
                unit = clean(sheet.cell(row, 3).value) or '-'
                unit_refs = [ref(row, 3)] if unit != '-' else []
                if '%' in header:
                    unit = '%'
                    unit_refs = [ref(header_row, col)]
                if annotation:
                    number, unit = annotation.groups()
                    unit_refs = [ref(row, col)]
                records.append({
                    'indicator_name': ''.join(name_parts) + suffix,
                    'indicator_data': number,
                    'indicator_unit': unit,
                    'end_date': date,
                })
                evidence.append({
                    'record_index': len(records) - 1,
                    'name_refs': [ref(p, 2) for p in parent_rows] + [ref(row, 2), ref(header_row, col)],
                    'data_ref': ref(row, col),
                    'unit_refs': unit_refs,
                    'date_refs': [ref(title_row, 2)] if date else [],
                })

    # 产品分类及县属子项分别绑定真实父项，不能让精煤继承原煤。
    product2 = {r: [10] for r in range(11, 16)}
    product2[12] = [10, 11]
    product2[14] = [10, 13]
    block('Sheet2', 2, 3, list(range(4, 10)) + list(range(11, 17)), product2)

    # 财政总收入与公共财政预算收入各自的组成项，避免同名税收错属。
    fiscal = {5: [4], 6: [4], 8: [7], 9: [7]}
    block('Sheet3', 1, 2, range(4, 10), fiscal, 3)
    block('Sheet3', 11, 12, range(14, 17), category_row=13)
    products3 = {19: [18], 21: [20]}
    block('Sheet3', 11, 12, range(18, 23), products3, 17)
    # B23既有章节序号又有真实数据，不能作为纯分类行跳过。
    block('Sheet3', 11, 12, [23])
    block('有', 1, 2, [3, 4, 5, 6, 7, 9])
    return records, evidence


def run(params, validation=False):
    try:
        md5 = params.get('fileMd5')
        if not md5:
            raise ValueError('缺少fileMd5')
        response = requests.get(DOWNLOAD_BASE + str(md5), timeout=60)
        response.raise_for_status()
        book = openpyxl.load_workbook(io.BytesIO(response.content), data_only=True)
        try:
            records, evidence = extract(book)
        finally:
            book.close()
        result = {'code': 0, 'msg': 'success', 'ExcelData': records}
        if validation:
            result['_dc_evidence'] = evidence
        return result
    except Exception as exc:
        result = {'code': 1, 'msg': str(exc), 'ExcelData': []}
        if validation:
            result['_dc_evidence'] = []
        return result


def main(params):
    return run(params)


def _dc_validate(params):
    return run(params, validation=True)
