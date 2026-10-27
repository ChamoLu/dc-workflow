# 数据源：信阳市浉河区人民政府
# 数据源路径代码：2082489640241225729
# 数据来源编码：01020003388
# 路径：信阳市浉河区人民政府\首页\政府信息公开\数据发布\数据月报（数据集可下载）
# 目标表：usrDFJJYXSJB
# 统计期从源表标题动态解析，不使用固定样例日期；GZK仅用于设计和审查，不参与运行时筛选。

import calendar
import datetime as dt
import io
import math
import re

import openpyxl
import requests


_DOWNLOAD_URL = "http://service-1.example.invalid/api/file/download/"


def _clean_text(value):
    if value is None:
        return ""
    text = str(value).replace("\r", "").replace("\n", "")
    text = text.replace("、", "").replace(",", "").replace("'", "")
    text = text.strip()
    text = re.sub(r"\s+", "", text)
    return text


def _is_invalid(value):
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip() in ("", "-")
    if isinstance(value, bool):
        return True
    if isinstance(value, float) and not math.isfinite(value):
        return True
    return False


def _parse_end_date(title):
    if not title:
        return ""
    text = str(title)
    match = re.search(r"(20\d{2})年\s*(\d{1,2})\s*月", text)
    if not match:
        match = re.search(r"(20\d{2})年\s*(\d{1,2})\s*[-至]\s*(\d{1,2})月", text)
        if match:
            year = int(match.group(1))
            month = int(match.group(3))
        else:
            return ""
    else:
        year = int(match.group(1))
        month = int(match.group(2))
    if month < 1 or month > 12:
        return ""
    return "%04d-%02d-%02d" % (year, month, calendar.monthrange(year, month)[1])


def _download(params):
    file_md5 = params.get("fileMd5")
    if not file_md5:
        return None, "missing fileMd5"
    response = requests.get(
        _DOWNLOAD_URL + str(file_md5),
        timeout=60,
    )
    response.raise_for_status()
    return response.content, ""


def _extract(params, with_evidence=False):
    raw, error = _download(params)
    if raw is None:
        return {"code": 1, "msg": error, "ExcelData": []}, []

    workbook = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    if "1-2月" in workbook.sheetnames:
        sheet = workbook["1-2月"]
    else:
        sheet = workbook[workbook.sheetnames[0]]

    table_title = sheet["A1"].value
    end_date = _parse_end_date(table_title)
    records = []
    evidence = []

    def add_record(row_number, column_number, header_ref, unit, data):
        name = _clean_text(sheet.cell(row=row_number, column=1).value)
        header = _clean_text(sheet[header_ref].value)
        if not name or not header or _is_invalid(data):
            return
        record = {
            "indicator_name": name + header,
            "indicator_data": data,
            "indicator_unit": unit if unit else "-",
            "end_date": end_date,
        }
        record_index = len(records)
        records.append(record)
        if with_evidence:
            unit_refs = []
            if column_number == 3:
                unit_refs = ["1-2月!B%d" % row_number] if unit else []
            else:
                unit_refs = ["1-2月!D3"]
            evidence.append({
                "record_index": record_index,
                "name_refs": ["1-2月!A%d" % row_number, "1-2月!%s" % header_ref],
                "data_ref": "1-2月!%s%d" % ("C" if column_number == 3 else "D", row_number),
                "unit_refs": unit_refs,
                "date_refs": ["1-2月!A1"] if end_date else [],
            })

    for row_number in range(4, 26):
        absolute_value = sheet.cell(row=row_number, column=3).value
        absolute_unit = _clean_text(sheet.cell(row=row_number, column=2).value) or "-"
        add_record(row_number, 3, "C3", absolute_unit, absolute_value)

        growth_value = sheet.cell(row=row_number, column=4).value
        add_record(row_number, 4, "D3", "%", growth_value)

    workbook.close()
    return {"code": 0, "msg": "success", "ExcelData": records}, evidence


def main(params):
    result, _ = _extract(params, with_evidence=False)
    return result


def _dc_validate(params):
    result, evidence = _extract(params, with_evidence=True)
    result["_dc_evidence"] = evidence
    return result
