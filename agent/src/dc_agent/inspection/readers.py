"""Bounded source inspection; preserve source cells without inferring business facts."""

from __future__ import annotations

import csv
import gzip
import hashlib
import io
import math
import zipfile
from datetime import date, datetime, time

MAX_BYTES = 64 * 1024 * 1024
MAX_ZIP_MEMBERS = 10000


class InspectionError(ValueError):
    """The source cannot be inspected completely within the configured bounds."""


def _value(value: object) -> object:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        raise InspectionError("Source contains a non-finite number")
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    raise InspectionError(f"Unsupported cell value type: {type(value).__name__}")


def _cell(ref: str, value: object, **extra: object) -> dict:
    return {"ref": ref, "value": _value(value), "python_type": type(value).__name__, **extra}


def _section(identifier: str, title: str, cells: list, merged: list) -> dict:
    return {
        "id": identifier,
        "title": title,
        "kind": "table" if cells else "empty",
        "cells": cells,
        "merged_ranges": merged,
    }


def _unwrap(data: bytes) -> bytes:
    for _ in range(3):
        if not data.startswith(b"\x1f\x8b"):
            return data
        try:
            with gzip.GzipFile(fileobj=io.BytesIO(data)) as source:
                data = source.read(MAX_BYTES + 1)
        except (OSError, EOFError) as exc:
            raise InspectionError("Invalid gzip source") from exc
        if len(data) > MAX_BYTES:
            raise InspectionError("Decompressed source exceeds byte limit")
    if data.startswith(b"\x1f\x8b"):
        raise InspectionError("Too many nested gzip wrappers")
    return data


def _check_xlsx_archive(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_MEMBERS:
                raise InspectionError("XLSX archive has too many members")
            if sum(item.file_size for item in entries) > MAX_BYTES:
                raise InspectionError("XLSX expanded archive exceeds byte limit")
            if any(item.flag_bits & 1 for item in entries):
                raise InspectionError("Encrypted XLSX archives are unsupported")
            if any(item.file_size > max(1, item.compress_size) * 1000 for item in entries):
                raise InspectionError("XLSX archive exceeds compression ratio limit")
            names = archive.namelist()
            if len(names) != len(set(names)):
                raise InspectionError("XLSX archive contains duplicate members")
            if "xl/workbook.xml" not in names or "[Content_Types].xml" not in names:
                raise InspectionError("ZIP source is not an XLSX workbook")
    except zipfile.BadZipFile as exc:
        raise InspectionError("Invalid ZIP/XLSX source") from exc


def _xlsx(data: bytes, max_cells: int) -> list:
    import openpyxl

    _check_xlsx_archive(data)
    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
    formulas = None
    try:
        # A formula without a cached result reads as None in data_only mode. Keep
        # its source position explicitly rather than claiming the cell is empty.
        formulas = openpyxl.load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
        sections = []
        visited = 0
        for sheet in workbook.worksheets:
            visited += sheet.max_row * sheet.max_column
            if visited > max_cells:
                raise InspectionError(
                    "Workbook grid exceeds max_cells; no partial inspection returned"
                )
            cells = []
            for row in sheet.iter_rows():
                for cell in row:
                    original = formulas[sheet.title][cell.coordinate]
                    if original.data_type == "f":
                        cells.append(
                            _cell(
                                f"{sheet.title}!{cell.coordinate}",
                                cell.value,
                                formula=str(original.value),
                                cache_missing=cell.value is None,
                            )
                        )
                    elif cell.value is not None:
                        cells.append(_cell(f"{sheet.title}!{cell.coordinate}", cell.value))
            sections.append(
                _section(
                    sheet.title, sheet.title, cells, [str(r) for r in sheet.merged_cells.ranges]
                )
            )
        return sections
    finally:
        workbook.close()
        if formulas is not None:
            formulas.close()


def _xls(data: bytes, max_cells: int) -> list:
    import xlrd
    from openpyxl.utils import get_column_letter

    workbook = xlrd.open_workbook(file_contents=data, formatting_info=True, on_demand=True)
    try:
        sections = []
        visited = 0
        for sheet in workbook.sheets():
            visited += sheet.nrows * sheet.ncols
            if visited > max_cells:
                raise InspectionError(
                    "Workbook grid exceeds max_cells; no partial inspection returned"
                )
            cells = []
            for row in range(sheet.nrows):
                for col in range(sheet.ncols):
                    cell = sheet.cell(row, col)
                    if cell.ctype not in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        cells.append(
                            _cell(
                                f"{sheet.name}!{get_column_letter(col + 1)}{row + 1}",
                                cell.value,
                                ctype=cell.ctype,
                            )
                        )
            merged = [
                f"{get_column_letter(c1 + 1)}{r1 + 1}:{get_column_letter(c2)}{r2}"
                for r1, r2, c1, c2 in sheet.merged_cells
            ]
            sections.append(_section(sheet.name, sheet.name, cells, merged))
        return sections
    finally:
        workbook.release_resources()


def _decode(data: bytes) -> str:
    encodings = (
        ("utf-16",) if data.startswith((b"\xff\xfe", b"\xfe\xff")) else ("utf-8-sig", "gb18030")
    )
    for encoding in encodings:
        try:
            value = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        if any(ord(char) < 32 and char not in "\t\n\r" for char in value):
            raise InspectionError("Text source contains binary control characters")
        return value
    raise InspectionError("Unsupported text encoding; expected UTF-8, BOM UTF-16, or GB18030")


def _text(data: bytes, file_type: str, max_cells: int) -> tuple[str, list]:
    content = _decode(data)
    if file_type in {"csv", "tsv"}:
        from openpyxl.utils import get_column_letter

        delimiter = "\t" if file_type == "tsv" else ","
        try:
            dialect = csv.Sniffer().sniff(content[:65536], delimiters=",\t;|")
            delimiter = dialect.delimiter
        except csv.Error:
            pass
        cells = []
        visited = 0
        for row_number, row in enumerate(
            csv.reader(io.StringIO(content, newline=""), delimiter=delimiter, strict=True), 1
        ):
            visited += max(1, len(row))
            if visited > max_cells:
                raise InspectionError("CSV exceeds max_cells; no partial inspection returned")
            for col, value in enumerate(row, 1):
                if value != "":
                    cells.append(_cell(f"csv!{get_column_letter(col)}{row_number}", value))
        return "csv", [_section("csv", "csv", cells, [])]
    if file_type not in {"markdown", "md", "html"}:
        raise InspectionError(f"Unsupported source format: {file_type}")
    lines = content.splitlines()
    if len(lines) > max_cells:
        raise InspectionError("Markdown exceeds max_cells; no partial inspection returned")
    section = _section(
        "markdown", "markdown", [_cell(f"L{i}", line) for i, line in enumerate(lines, 1)], []
    )
    section["kind"] = "markdown" if lines else "empty"
    return "markdown", [section]


def inspect_bytes(data: bytes, file_type: str, *, max_cells: int = 50000) -> dict:
    """Inspect every source section or raise; byte magic takes precedence for Excel."""
    if not isinstance(data, bytes):
        raise InspectionError("Source must be bytes")
    if not isinstance(max_cells, int) or isinstance(max_cells, bool) or max_cells < 1:
        raise InspectionError("max_cells must be a positive integer")
    if len(data) > MAX_BYTES:
        raise InspectionError("Source exceeds byte limit")
    source = _unwrap(data)
    try:
        if source.startswith(b"PK"):
            format_name, sections = "xlsx", _xlsx(source, max_cells)
        elif source.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            format_name, sections = "xls", _xls(source, max_cells)
        else:
            format_name, sections = _text(source, file_type.strip().lower().lstrip("."), max_cells)
    except InspectionError:
        raise
    except Exception as exc:
        raise InspectionError(f"Cannot inspect source ({type(exc).__name__}): {exc}") from exc
    return {
        "format": format_name,
        "sha256": hashlib.sha256(data).hexdigest(),
        "size": len(data),
        "coverage_complete": True,
        "sections": sections,
    }
