"""Read fresh workbook snapshots; never infer tasks or sheets from similar names."""

import hashlib
import io
import math
from contextlib import contextmanager
from datetime import date, datetime, time

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..config import Settings
from ..domain import AgentError, Task


def _value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("Non-finite reference cell")
    return value


def _path_code(value) -> str | None:
    if isinstance(value, str):
        # Do not normalize identifiers or remove significant leading zeros/spaces.
        return value if value else None
    # Excel cannot retain more than 15 significant numeric digits. Do not guess
    # an identifier from a rounded long number; the task may proceed without a sheet.
    if (
        isinstance(value, (int, float))
        and not isinstance(value, bool)
        and math.isfinite(value)
        and value == int(value)
        and abs(value) < 10**15
    ):
        return str(int(value))
    return None


class WorkbookRuleContext:
    def __init__(self, settings: Settings):
        self.settings = settings

    @contextmanager
    def _open(self, filename: str):
        path = self.settings.resolved_rules() / filename
        try:
            with path.open("rb") as stream:
                raw = stream.read(self.settings.max_rule_bytes + 1)
        except OSError as exc:
            raise AgentError(
                "REFERENCE_UNAVAILABLE",
                f"Cannot read rule workbook: {path.name}",
                stage="resolve_rule_context",
                retryable=True,
            ) from exc
        if len(raw) > self.settings.max_rule_bytes:
            raise AgentError(
                "REFERENCE_TOO_LARGE",
                f"Rule workbook exceeds byte limit: {path.name}",
                stage="resolve_rule_context",
            )
        workbook = None
        try:
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            yield workbook, {"file": path.name, "sha256": hashlib.sha256(raw).hexdigest()}
        except AgentError:
            raise
        except Exception as exc:
            raise AgentError(
                "REFERENCE_FORMAT",
                f"Invalid rule workbook structure: {path.name}",
                stage="resolve_rule_context",
            ) from exc
        finally:
            if workbook is not None:
                workbook.close()

    def _task_info(self, task: Task) -> dict:
        matches, has_task_sheet = [], False
        with self._open(self.settings.todo_filename) as (workbook, origin):
            # A task table is identified by its C/E headers, not a fixed sheet name.
            # Multiple task sheets participate in the same uniqueness check.
            for sheet in workbook:
                rows = sheet.iter_rows()
                headers = next(rows, ())
                if len(headers) < 5 or (
                    str(headers[2].value or "").strip() != "DataSourcePathCode"
                    or str(headers[4].value or "").strip() != "pageTitle"
                ):
                    continue
                has_task_sheet = True
                for row in rows:
                    if row[4].data_type == "f":
                        raise AgentError(
                            "REFERENCE_FORMAT",
                            "todo pageTitle must be literal text, not a formula",
                            stage="resolve_rule_context",
                        )
                    if row[4].value != task.pageTitle:
                        continue
                    fields = {}
                    for column, (header, cell) in enumerate(zip(headers, row), 1):
                        key = str(header.value).strip() if header.value is not None else ""
                        if not key:
                            key = get_column_letter(column)
                        if key in fields:
                            key = f"{get_column_letter(column)}:{key}"
                        fields[key] = _value(cell.value)
                    matches.append(
                        {
                            "source": {**origin, "sheet": sheet.title, "row": row[4].row},
                            "fields": fields,
                            "data_source_path_code": _path_code(row[2].value),
                        }
                    )
        if not has_task_sheet:
            raise AgentError(
                "REFERENCE_FORMAT",
                "todo requires C=DataSourcePathCode and E=pageTitle headers",
                stage="resolve_rule_context",
            )
        if not matches:
            raise AgentError(
                "todoFile_no_pageTitle",
                "No exact pageTitle in todo workbook",
                stage="resolve_rule_context",
            )
        if len(matches) > 1:
            raise AgentError(
                "pageTitle_redundancy",
                "Multiple task rows have the exact pageTitle",
                stage="resolve_rule_context",
                details={"matching_rows": [m["source"] for m in matches]},
            )
        return matches[0]

    def load(self, task: Task) -> dict:
        info = self._task_info(task)
        references = []
        code = info["data_source_path_code"]
        if code:
            is_html = task.kind == "python_html"
            filename = (
                self.settings.html_reference_filename
                if is_html
                else self.settings.indicator_reference_filename
            )
            with self._open(filename) as (workbook, origin):
                if code in workbook.sheetnames:
                    sheet = workbook[code]
                    # Both reference formats may be headerless. Preserve every
                    # populated cell, its position and the first row for the model.
                    cells = [
                        {"address": cell.coordinate, "value": _value(cell.value)}
                        for row in sheet.iter_rows()
                        for cell in row
                        if cell.value is not None
                    ]
                    if cells:
                        references.append(
                            {
                                **origin,
                                "sheet": sheet.title,
                                "role": "html_name_regex"
                                if is_html
                                else "indicator_name_fragments",
                                "cells": cells,
                            }
                        )
        # A missing sheet is an ordinary optional-reference branch: no absence
        # marker, warning, error or model explanation is added.
        return {"task_info": info, "references": references}
